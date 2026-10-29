"""State court cases found online: a step of the daily update (update), which puts what both readers verify on the
dashboard and the rest on the To check list. A dry run that changes nothing but its own folder:

    .venv\\Scripts\\python app\\state_search.py        (writes State cases\\Search trial\\<date>\\results.xlsx)

CourtListener misses most state trial-court rulings, so this looks for state cases itself, then has them verified:

1. Find. a) Claude Code with only its web search tool runs SEARCHES and lists the state cases it finds, with sources:
   news and law-firm searches, every state by name, and the courts' own sites as the search engine has them.
   (Every state court's posted opinions on CourtListener are already searched by the daily update: STATE_QUERY.)
   b) The NJ courts' own opinion lists (newest page of each) - an opinion whose caption names an insurer and whose
   PDF mentions the No Surprises Act. c) The NJ multicounty litigation page, for the day the NSA cases are grouped.
2. Skip what the dashboard already has (both parties' names match a case on it).
3. Fetch each source, only where the site's robots.txt allows it; a site that blocks scripts or shows a bot check is
   skipped (never worked around), and the next source for the same case is tried; then CourtListener's copy of the
   court's opinion (cl_copy), then a search for other free copies (more_sources).
4. Verify with the same two readers (Sonnet and Opus). A court document is read exactly as a file saved in
   State cases\\New would be (state_cases.check_file); a news or law-firm page is read by read_news. Both readers must
   agree and their quote must be in the text. If every page blocks scripts, both readers check it through web
   search (read_by_search). A lead not verified after MAX_TRIES runs goes to To check.
"""
import datetime as dt
import difflib
import html
import io
import json
import re
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import daily_update as du  # noqa: E402
import data  # noqa: E402
import insurers  # noqa: E402
import reader  # noqa: E402
import state_cases  # noqa: E402
import find_updates as fu  # noqa: E402  (app\tools, put on the path by daily_update)
from find_updates import split_parties, strip_tags  # noqa: E402

UA = "LitigationTracker/1.0 (internal research; state court cases)"
PAUSE = 1.5  # seconds between requests to the same site

# ---------------- 1a. web search ----------------

SEARCHES = [
    '"No Surprises Act" IDR awards lawsuit "Superior Court" provider insurer',
    '"No Surprises Act" IDR award "Supreme Court" New York County OR Nassau OR Suffolk OR Kings provider sues',
    '"No Surprises Act" IDR awards New York state court decision provider insurer',
    '"No Surprises Act" "Appellate Division" IDR award decision',
    '"No Surprises Act" IDR awards Florida circuit court provider lawsuit insurer',
    '"No Surprises Act" IDR awards Texas state district court provider lawsuit',
    '"No Surprises Act" IDR awards Georgia OR California OR Colorado OR Illinois state court lawsuit',
    '"No Surprises Act" IDR awards state court ruling dismissed OR granted',
    'unpaid "IDR awards" lawsuit state court New Jersey OR New York',
    '"No Surprises Act" "Prompt Pay" IDR determinations state law claims court decision',
    'beckerspayer OR healthcaredive OR law360 "No Surprises" awards "state court" provider sues',
    'jdsupra OR mondaq OR natlawreview "No Surprises Act" state court IDR award enforcement decision',
    'New York and Florida courts reject provider suits to enforce NSA IDR awards',
    # courts' own sites, as the search engine has them (nycourts.gov shows scripts a bot check, so it is never fetched)
    'site:nycourts.gov "No Surprises Act" IDR award decision',
    '"No Surprises Act" IDR site:flcourts.gov OR site:txcourts.gov OR site:courts.ca.gov OR site:coloradojudicial.gov '
    'OR site:illinoiscourts.gov OR site:pacourts.us OR site:mass.gov',
    'filetype:pdf "No Surprises Act" "independent dispute resolution" "Supreme Court of the State of New York" OR '
    '"Superior Court" OR "Circuit Court" OR "District Court of" decision provider insurer',
]
_NAMED = {"NJ", "NY", "FL", "TX", "GA", "CA", "CO", "IL"}  # states the searches above already name
_REST = [n for c, n in data.STATE_NAMES.items() if c not in _NAMED]
SEARCHES += ['"No Surprises Act" IDR awards state court lawsuit provider insurer ' + " OR ".join(_REST[i:i + 7])
             for i in range(0, len(_REST), 7)]  # every other state, seven to a search
GROUP = 4  # searches per Claude call: one long call reports fewer cases than several short ones
FIND_RULES = """You find state-court lawsuits for a litigation tracker. The tracker follows suits in which a medical
provider asks a court to confirm, enforce or collect on No Surprises Act IDR (independent dispute resolution) awards
against an insurer or health plan, filed in a STATE court (not a federal court), in any US state, from 2022 on.

Use web search only. Run every search listed in the task, plus follow-up searches that help name the case, court and
date, or find a free copy of the court's own document. Report each distinct state-court case you find: a ruling in
one, or a new case being filed. Skip federal cases (U.S. District Courts, Circuits) and suits brought by insurers.
Each case must be one named case ("Plaintiff v. Defendant"); skip mentions of unnamed groups of cases.

For each case list up to four source URLs, best first: the court's own document (a PDF on a court website) before a
law-firm write-up before a news article. Prefer pages that are free to read (no paywall). Copy the sources' own words
for the case in 'evidence'. Do not guess: leave a field empty when the sources do not say it."""
FIND_SCHEMA = {"type": "object", "required": ["cases"], "properties": {"cases": {"type": "array", "items": {
    "type": "object", "required": ["case", "court", "state", "date", "kind", "sources", "evidence"],
    "properties": {"case": {"type": "string"}, "court": {"type": "string"}, "state": {"type": "string"},
                   "date": {"type": "string"}, "kind": {"type": "string", "enum": ["ruling", "new case", "unknown"]},
                   "sources": {"type": "array", "items": {"type": "string"}}, "evidence": {"type": "string"}}}}}}


SOURCES_RULES = """You find free, readable copies of what is written about one state-court lawsuit. Use web search only.
Look for the court's own document first (a PDF on a court website or on a law school or law-firm site), then law-firm
write-ups (JD Supra, Mondaq, National Law Review, firm blogs), then news. Skip pages behind a paywall, the sites
listed as unreadable, and paid case-data services (UniCourt, Trellis, PacerMonitor, Docket Alarm, Westlaw, Lexis). Only list pages that are about this exact case. Copy their own words for it in 'evidence'."""
SOURCES_SCHEMA = {"type": "object", "required": ["sources", "evidence"], "properties": {
    "sources": {"type": "array", "items": {"type": "string"}}, "evidence": {"type": "string"}}}


def web_leads(searches=SEARCHES, exe=None):
    """Leads from Claude Code's web search, GROUP searches per call, the same case reported twice kept once."""
    leads = []
    for i in range(0, len(searches), GROUP):
        task = "Run these searches:\n" + "\n".join(f"- {s}" for s in searches[i:i + GROUP])
        for c in search_call(FIND_RULES, FIND_SCHEMA, task, exe)["cases"]:
            add_lead(leads, {**c, "foundBy": "web search"})
    return leads


def add_lead(leads, c):
    """Add c to leads, or, when it is a case already there (same parties, or the same article), add its sources to
    that one. Returns the lead that was new, or None."""
    twin = next((x for x in leads if names(x["case"]) == names(c["case"]) and all(names(c["case"]))
                 or set(x["sources"]) & set(c["sources"])), None)
    if twin:
        twin["sources"] += [u for u in c["sources"] if u not in twin["sources"]]
        return None
    leads.append(c)
    return c


def more_sources(lead, unreadable, exe=None):
    """Up to four more source URLs for one lead whose sources could not be read."""
    hosts = sorted({urllib.parse.urlsplit(u).netloc for u in unreadable})
    task = (f"CASE: {lead['case']}\nCOURT: {lead['court']} ({lead['state']})\nWHAT IS KNOWN: {lead['evidence'][:600]}\n"
            f"UNREADABLE SITES: {', '.join(hosts) or 'none'}")
    return search_call(SOURCES_RULES, SOURCES_SCHEMA, task, exe)["sources"][:4]


def search_call(rules, schema, task, exe=None, model=reader.L1):
    """One Claude Code call with the WebSearch tool only (no reading or fetching of pages, no files, no commands),
    answering in schema, so a web page's text can only ever be data."""
    cmd = [exe or reader.claude_exe(), "-p", "--model", model, "--tools", "WebSearch", "--allowedTools", "WebSearch",
           "--safe-mode", "--strict-mcp-config", "--no-session-persistence", "--permission-prompts", "none",
           "--system-prompt", rules, "--json-schema", json.dumps(schema), "--output-format", "json"]
    try:
        p = subprocess.run(cmd, input=task, capture_output=True, text=True, encoding="utf-8", timeout=reader.TIMEOUT,
                           cwd=tempfile.gettempdir())
        if p.returncode:
            raise reader.ReaderError(f"Claude stopped ({(p.stderr or p.stdout).strip()[-600:]})")
        out = json.loads(p.stdout)
        return out.get("structured_output") or json.loads(out["result"])
    except (OSError, ValueError, KeyError, subprocess.SubprocessError) as e:
        raise reader.ReaderError(f"Claude could not be run ({e})") from None


# ---------------- 1b, 1c. NJ courts ----------------

NJ = "https://www.njcourts.gov"
NJ_LISTS = ("unpublished-trial", "published-trial", "unpublished-appellate", "published-appellate", "supreme")
CAPTION = re.compile("|".join(p for p, _, k, _ in insurers.RULES if k in (insurers.INS, insurers.BLUE, insurers.NET,
                                                                          insurers.PLAN) and "^" not in p)
                     + r"|insurance|\bins\.|health plan|benefit", re.I)
NSA = re.compile(r"no surprises act|independent dispute resolution|300gg-11[1-9]", re.I)


def nj_rows(page):
    """(pdf url, caption text) for each opinion on one NJ opinion list page."""
    main = page[page.find("<main"):]
    parts = re.split(r'(href="[^"]+\.pdf[^"]*")', main)
    for i in range(1, len(parts) - 1, 2):
        yield (urllib.parse.urljoin(NJ, re.search(r'href="([^"]+)"', parts[i])[1]),
               # parts[i + 1] starts inside the link's tag and ends inside the next one's
               re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>|<[^>]*$", " ", "<a" + parts[i + 1]))).strip()[:200])


def nj_leads(web, pages=1, seen=None):
    """Opinions on the newest pages of NJ's opinion lists whose caption names an insurer and whose PDF mentions the No
    Surprises Act (a word search, no AI), plus a lead for an NSA entry on the multicounty litigation page. seen: the
    PDFs checked before (added to, so each is downloaded once)."""
    leads, seen = [], [] if seen is None else seen
    for lst in NJ_LISTS:
        for n in range(pages):
            page = web.text(f"{NJ}/attorneys/opinions/{lst}" + (f"?page={n}" if n else ""))
            for url, caption in nj_rows(page):
                if url in seen or not CAPTION.search(caption):
                    continue
                pdf = web.get(url)
                seen.append(url)
                if NSA.search(pdf_words(pdf)):
                    leads.append({"case": caption, "court": "New Jersey courts", "state": "NJ", "date": "",
                                  "kind": "ruling", "sources": [url], "evidence": caption, "foundBy": f"NJ courts ({lst})"})
    page = web.text(f"{NJ}/attorneys/multicounty-litigation")
    for h, t in re.findall(r'<a[^>]+href="([^"]*multicounty-litigation/[^"]+)"[^>]*>(.*?)</a>', page, re.S):
        if re.search(r"surprise|dispute resolution|\bIDR\b", t + h, re.I):
            leads.append({"case": strip_tags(t), "court": "New Jersey Superior Court (multicounty litigation)",
                          "state": "NJ", "date": "", "kind": "unknown", "sources": [urllib.parse.urljoin(NJ, h)],
                          "evidence": "NJ grouped cases: " + strip_tags(t), "foundBy": "NJ multicounty page"})
    return leads


def pdf_words(pdf):
    try:
        import pypdf
        return " ".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(pdf)).pages[:40])
    except Exception:  # a damaged or scanned PDF: no words to search
        return ""


# ---------------- 3. fetching, politely ----------------

class Blocked(Exception):
    """The site does not allow scripts here (robots.txt), refused, showed a bot check or a paywall, or did not answer."""


# Paid case-data services (only free, public sources are used) and sites whose terms forbid automated access.
NOT_USED = re.compile(r"(^|\.)(unicourt\.com|trellis\.law|westlaw\.com|lexis\.com|lexisnexis\.com|casetext\.com|"
                      r"pacermonitor\.com|docketalarm\.com|justia\.com|leagle\.com|scholar\.google\.com|"
                      r"bloomberglaw\.com|courtlink\.lexisnexis\.com)$|^iapps\.courts\.state\.ny\.us$", re.I)
BOT_CHECK = re.compile(r"captcha|crowdsec|verify you are (a )?human|are you a robot|just a moment\.\.\.|cf-chl|"
                       r"enable javascript and cookies|access denied|subscribe to (continue|read)", re.I)


class Web:
    def __init__(self):
        self.robots, self.last = {}, {}

    def allowed(self, url):
        origin = "{0.scheme}://{0.netloc}".format(urllib.parse.urlsplit(url))
        if origin not in self.robots:
            rp = urllib.robotparser.RobotFileParser()
            try:
                with urllib.request.urlopen(urllib.request.Request(origin + "/robots.txt", headers={"User-Agent": UA}),
                                            timeout=30) as r:
                    rp.parse(r.read(500_000).decode("utf-8", "replace").splitlines())
            except urllib.error.HTTPError as e:  # the standard (RFC 9309): no robots.txt (4xx) means no rules,
                rp.allow_all = 400 <= e.code < 500  # a server error means stay away
                rp.disallow_all = not rp.allow_all
            except (OSError, ValueError):
                rp.disallow_all = True
            self.robots[origin] = rp
        return self.robots[origin].can_fetch(UA, url)

    def get(self, url):
        """The bytes at url. Raises Blocked."""
        if NOT_USED.search(urllib.parse.urlsplit(url).netloc):
            raise Blocked("a paid service, or a site whose terms forbid scripts: not used")
        if not self.allowed(url):
            raise Blocked("the site's robots.txt does not allow scripts here")
        host = urllib.parse.urlsplit(url).netloc
        time.sleep(max(0, self.last.get(host, 0) + PAUSE - time.time()))
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA}), timeout=60) as r:
                body = r.read(30_000_000)
        except urllib.error.HTTPError as e:
            raise Blocked(f"the site refused (HTTP {e.code})") from None
        except (OSError, ValueError) as e:
            raise Blocked(f"the site did not answer ({e})") from None
        finally:
            self.last[host] = time.time()
        if body[:5] != b"%PDF-" and BOT_CHECK.search(body[:20000].decode("utf-8", "replace")):
            raise Blocked("the site showed a bot check or a paywall")
        return body

    def text(self, url):
        return self.get(url).decode("utf-8", "replace")


def page_text(body):
    raw = body.decode("utf-8", "replace")
    raw = re.sub(r"(?is)<(script|style|nav|header|footer|aside|form)\b.*?</\1>", " ", raw)
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", raw))).strip()


# ---------------- 4. verifying a news or law-firm page ----------------

NEWS_RULES = """You read one web page (a news article, a law-firm write-up or a court notice) for a litigation tracker
of lawsuits in which a medical provider asks a STATE court to confirm, enforce or collect on No Surprises Act IDR
(independent dispute resolution) awards against an insurer or health plan. Answer about the case named in LEAD only,
using only the page's text:
- reported: true only if the page reports that case, in a state court (not a federal court), and it is such a suit.
- case: "First plaintiff v. First defendant" as the page names them. Never name a patient.
- court: the court as the page names it, with the county if given.
- state: the two-letter code of the court's state.
- date: YYYY-MM-DD of the ruling (for a ruling) or the filing (for a new case). Only if the page gives the full date
  including the year, or the year is certain from the page's own date line; otherwise "".
- kind: "ruling" if the page reports a court decision in the case, "new case" if only that it was filed.
- result: for a ruling, "Win" (the provider's claim to enforce or collect on the awards was allowed or confirmed),
  "Loss" (it was dismissed or rejected) or "Pending" (the ruling does not decide that). "" for a new case.
- holding: one plain-English sentence on what the court decided or what the suit asks for.
- quote: copy word for word the one sentence of the page that best shows the court and what happened."""
NEWS_SCHEMA = {"type": "object", "required": ["reported", "case", "court", "state", "date", "kind", "result", "holding", "quote"],
               "properties": {"reported": {"type": "boolean"}, **{k: {"type": "string"} for k in
                              ("case", "court", "state", "date", "kind", "result", "holding", "quote")}}}


def read_news(text, lead, ask=None):
    """Two readers, as reader.identify: Sonnet and Opus each read the page on their own; they must agree on reported,
    state, kind, date and result, and both quotes must be in the page. Up to reader.ROUNDS rounds."""
    page = reader.clean(text)
    if len(page.split()) < 60:
        return {"agreed": False, "note": "The page has no readable text."}
    prompt = f"LEAD: {lead['case']} ({lead.get('court', '')})\n\nPAGE TEXT:\n\n{page[:60000]}"
    ask = ask or (lambda model, system, p: reader.ask_claude(model, system, p, schema=NEWS_SCHEMA))
    text_n = reader.norm(page[:60000])
    found = lambda a: len(reader.norm(a["quote"])) > 15 and reader.norm(a["quote"]) in text_n  # noqa: E731
    key = lambda a: (a["reported"], a["state"].upper(), a["kind"], a["date"], a["result"])  # noqa: E731
    a, b = ask(reader.L1, NEWS_RULES, prompt), ask(reader.L2, NEWS_RULES, prompt)
    for rnd in range(1, reader.ROUNDS + 1):
        if found(a) and found(b) and key(a) == key(b):
            return {**a, "state": a["state"].upper(), "agreed": True, "note": ""}
        if rnd == reader.ROUNDS:
            return {**a, "agreed": False, "note": "The two readers did not agree on what this page reports."}
        again = lambda other, ok: ("\n\nYour quote was not found word for word in the page. Copy the sentence exactly."  # noqa: E731
                                   if not ok else f"\n\nAnother reader answered: {json.dumps(other)}. Read the page again "
                                   "and give your answer. Change it only if the text supports the change.")
        a, b = ask(reader.L1, NEWS_RULES, prompt + again(b, found(a))), ask(reader.L2, NEWS_RULES, prompt + again(a, found(b)))


SEARCH_READ_RULES = NEWS_RULES.replace("You read one web page (a news article, a law-firm write-up or a court notice)",
                                       "You check what is published about one lawsuit, using web search only,").replace(
    "using only the page's text:", "using only what the search results say:").replace(
    "copy word for word the one sentence of the page", "copy word for word the one sentence of a search result") + """
- quotes: that sentence first, then up to two more sentences about this case copied word for word from the search
  results (for a news report, its opening sentence).
- source: the URL of the page the first sentence comes from."""
SEARCH_READ_SCHEMA = {**NEWS_SCHEMA, "required": NEWS_SCHEMA["required"] + ["quotes", "source"],
                      "properties": {**NEWS_SCHEMA["properties"], "source": {"type": "string"},
                                     "quotes": {"type": "array", "items": {"type": "string"}}}}
SAME_QUOTE = 40  # letters both readers' quotes must share word for word


def read_by_search(lead, ask=None):
    """For a lead whose pages all block scripts: Sonnet and Opus each search the web on their own. They must agree on
    reported, state, kind, date and result, and one quote of each must share SAME_QUOTE letters word for word - two
    separate readers only copy the same words from the same real page. When they differ, each is shown the other's
    answer WITHOUT its quotes (so the quotes stay their own) and answers again; up to reader.ROUNDS rounds."""
    task = (f"LEAD: {lead['case']} ({lead.get('court', '')}, {lead.get('state', '')})\n"
            f"WHAT A SEARCH SAID: {lead.get('evidence', '')[:600]}")
    ask = ask or (lambda model, t=task: search_call(SEARCH_READ_RULES, SEARCH_READ_SCHEMA, t, model=model))
    key = lambda x: (x["reported"], x["state"].upper(), x["kind"], x["result"])  # noqa: E731
    shown = lambda x: {k: x[k] for k in ("reported", "case", "court", "state", "date", "kind", "result", "source")}  # noqa: E731
    a, b = ask(reader.L1), ask(reader.L2)
    for rnd in range(1, reader.ROUNDS + 1):
        # a date only one reader gives counts as no date (a person then checks it); two different dates disagree
        if (key(a) == key(b) and (a["date"] == b["date"] or not (a["date"] and b["date"])) and same_words(a, b)
                and a["source"].startswith("http")):
            return {**a, "state": a["state"].upper(), "date": a["date"] if a["date"] == b["date"] else "",
                    "agreed": True, "note": ""}
        if rnd == reader.ROUNDS:
            return {**a, "agreed": False, "note": "The two readers' web searches did not agree on what was reported."}
        again = lambda other: (f"{task}\n\nAnother reader answered: {json.dumps(shown(other))}. Search again and give "  # noqa: E731
                               "your answer. Change it only if the search results support the change.")
        a, b = ask(reader.L1, again(b)), ask(reader.L2, again(a))


def same_words(a, b):
    """True when one of a's quotes and one of b's share SAME_QUOTE letters word for word."""
    for x in [a["quote"]] + a.get("quotes", []):
        for y in [b["quote"]] + b.get("quotes", []):
            qa, qb = reader.norm(x), reader.norm(y)
            if difflib.SequenceMatcher(None, qa, qb, autojunk=False).find_longest_match(0, len(qa), 0, len(qb)).size >= SAME_QUOTE:
                return True
    return False


def news_result(res, url, how="news report"):
    """(what happened, entries for found.json, why it needs a person or "") for an agreed news reading."""
    if not res["reported"]:
        return "to check", {}, "Both readers say this page does not report a state-court suit to enforce IDR awards."
    if res["state"] not in data.STATE_NAMES:
        return "to check", {}, "The sources do not say which state's court it is in."
    if not re.fullmatch(r"\d{4}-\d\d-\d\d", res["date"]):
        return "to check", {}, "The sources do not give the full date (with the year), so a person checks it."
    case = data.mask(res["case"])
    row = {"case": case, "court": res["court"], "level": "State Trial", "circuit": "", "state": res["state"],
           "docket": "", "payer": data.mask(split_parties(case)[1]), "url": url, "source": f"State court search ({how})"}
    if res["kind"] == "ruling" and res["result"] in ("Win", "Loss"):
        return (f"{res['result']}: shown as a ruling (from a {how})",
                {"decisions": [{**row, "date": res["date"], "result": res["result"], "final": False,
                                "holding": data.mask(res["holding"]), "fromNews": how}]}, "")
    return (f"shown as pending (from a {how})",
            {"pending": [{**row, "filed": res["date"], "fromNews": how,
                          "status": "Reported in the news; both readers confirmed it is a state suit to enforce IDR awards."}]}, "")


# ---------------- 2. what the dashboard already has ----------------

def names(case):
    """The first distinctive word of each party, e.g. ("true", "anthem") for True North Medical v. Anthem Inc."""
    return tuple(next((w for w in re.findall(r"[a-z]{3,}", side.lower()) if w not in du.NOT_NAMES), "")
                 for side in split_parties(case))


def known_case(case, cases):
    """The dashboard case with both parties' names, or None."""
    a, b = (names(case) + ("", ""))[:2]
    if not a or not b:
        return None
    return next((c for c in cases if a in c.lower() and b in c.lower()), None)


# ---------------- the run ----------------

def verify(lead, web, system, folder, n, hunt=True, cl=None):
    """Try the lead's sources in turn; the first readable one is read by both readers. If none can be used, the
    court's opinion is looked for on CourtListener (cl = (api, token)), then one targeted search (more_sources) looks
    for free copies elsewhere and those are tried too. If every page blocks scripts, the two readers check the case
    through web search (read_by_search). Returns {what, why, adds, readFrom, tried, details}."""
    res = _verify(lead["sources"][:4], lead, web, system, folder, n)
    if res["readFrom"] or not hunt:
        return res
    copy = cl_copy(lead, *cl) if cl else None
    if copy:
        return {**_read(copy[1], copy[0], lead, system, True), "tried": res["tried"]}
    extra = [u for u in more_sources(lead, lead["sources"]) if u not in lead["sources"]]
    again = _verify(extra, lead, web, system, folder, n)
    again["tried"] = res["tried"] + again["tried"]
    if again["readFrom"]:
        return again
    info = read_by_search(lead)  # every page blocks scripts: the readers check it through web search instead
    if not info["agreed"]:
        return {**again, "why": again["why"] + " " + info["note"], "details": info}
    what, adds, why = news_result(info, info["source"], "web search")
    return {"what": what, "why": why, "adds": adds, "readFrom": info["source"] + " (via web search)",
            "tried": again["tried"], "details": info}


def cl_copy(lead, api, token):
    """(url, text) of a state court's opinion in this case from CourtListener (both parties in its case name), or
    None. Courts whose own sites block scripts (nycourts.gov) are often there."""
    a, b = (names(lead["case"]) + ("", ""))[:2]
    if not (a and b and token):
        return None
    q = urllib.parse.urlencode({"type": "o", "q": f"caseName:({a} AND {b})", "order_by": "dateFiled desc"})
    try:
        results = api.get(f"{fu.CL}/api/rest/v4/search/?{q}", {"Authorization": "Token " + token}, cl=True)["results"]
    except (RuntimeError, KeyError):
        return None
    for r in results[:5]:
        c = fu.candidate_from_cl(r, "o")
        if c and c["level"].startswith("State"):
            text = du.order_text(api, token, {"fetch": [("cl_opinion", o["id"]) for o in r.get("opinions") or []]})
            if text and du.same_case(text, lead["case"]):
                return c["url"], text
    return None


def _verify(urls, lead, web, system, folder, n):
    tried = []
    for url in urls:
        try:
            body = web.get(url)
        except Blocked as e:
            tried.append(f"{url} - {e}")
            continue
        is_pdf = body[:5] == b"%PDF-"
        court_doc = is_pdf or re.search(r"court|\.gov/", urllib.parse.urlsplit(url).netloc + urllib.parse.urlsplit(url).path, re.I)
        f = folder / f"{n:02d} {re.sub(r'[^A-Za-z0-9.-]+', '_', urllib.parse.urlsplit(url).netloc)}{'.pdf' if is_pdf else '.html'}"
        f.write_bytes(body)
        text = du.file_text(f) if is_pdf else page_text(body)
        if not du.same_case(text, lead["case"]):
            tried.append(f"{url} - does not name the parties")
            continue
        return {**_read(text, url, lead, system, court_doc), "tried": tried}
    return {"what": "not verified", "why": "No source could be read (blocked, paywalled or another case).", "adds": {},
            "readFrom": "", "tried": tried, "details": {}}


def _read(text, url, lead, system, court_doc):
    """Both readers on one source's text: a court document as State cases\\New would read it, else as a news page."""
    try:
        if court_doc:
            what, adds, why, info = state_cases.check_file(text, system, lambda case: split_parties(case)[1])
            for k in adds:  # the source is online: link it
                for x in adds[k]:
                    x["url"] = url
        else:
            info = read_news(text, lead)
            what, adds, why = news_result(info, url) if info["agreed"] else ("to check", {}, info["note"])
    except reader.Declined:
        what, adds, why, info = "to check", {}, "Claude declined to read this source.", {}
    return {"what": what, "why": why, "adds": adds, "readFrom": url, "details": info}


MAX_TRIES = 5  # runs a lead is tried before it goes to To check for a person
MAX_LEADS = 10  # open leads tried per run; the rest wait for the next run


def search(leads, known, system, done, say=print, folder=None, nj_seen=None):
    """Find new leads (web search, NJ courts) into leads, then try up to MAX_LEADS open ones: skip a case on the
    dashboard (known: case names), else verify. A lead stays open until it is verified, goes to To check or is on the
    dashboard; after MAX_TRIES runs it goes to To check. done(lead, res) is called after each lead. Claude stopping
    (reader.ReaderError) ends the run; the rest are tried next run."""
    today = dt.date.today().isoformat()
    web, cl = Web(), (fu.Api(), fu.user_env("COURTLISTENER_TOKEN") or "")
    folder = Path(folder or tempfile.mkdtemp(prefix="state_search_"))
    new = []
    try:
        new += web_leads()
    except reader.ReaderError as e:
        say(f"  web search stopped: {e}. Leads found earlier are still tried.")
    try:
        new += nj_leads(web, 1, nj_seen)
    except Blocked as e:
        say(f"  NJ courts skipped: {e}")
    for c in new:
        c = {**c, "case": data.mask(c["case"]), "evidence": data.mask(c.get("evidence", ""))}
        if add_lead(leads, {**c, "status": "open", "tries": 0, "firstSeen": today}):
            say(f"  new lead: {c['case'][:90]}")
    todo = [x for x in leads if x["status"] == "open"]
    say(f"  {len(todo)} open leads; trying {min(len(todo), MAX_LEADS)}")
    for n, lead in enumerate(todo[:MAX_LEADS], 1):
        say(f"  [{n}] {lead['case'][:90]}")
        same = known_case(lead["case"], known)
        if same:
            res = {"what": "already known", "why": same, "adds": {}, "readFrom": "", "tried": [], "details": {}}
        else:
            res = verify(lead, web, system, folder, n, cl=cl)
            got = (res["adds"].get("decisions") or res["adds"].get("pending") or [{}])[0]
            if got and known_case(got["case"], known):
                res.update(what="already known", why=known_case(got["case"], known), adds={})
        if res["what"] == "not verified" and lead["tries"] + 1 >= MAX_TRIES:
            res["what"], res["why"] = "to check", f"Not verified after {MAX_TRIES} runs. " + res["why"]
        lead.update(status="open" if res["what"] == "not verified" else "done", tries=lead["tries"] + 1,
                    lastTried=today, outcome=res["what"])
        say(f"    -> {res['what']}" + (f" ({res['why'][:200]})" if res["why"] else ""))
        done(lead, res)


def state_code(s):
    """"NY" for "NY", "ny" or "New York"; "" if not a state."""
    s = (s or "").strip()
    return s.upper() if s.upper() in data.STATE_NAMES else {v.lower(): k for k, v in data.STATE_NAMES.items()}.get(s.lower(), "")


def update(report, state, found, store, asks, say=print, progress=lambda: None):
    """The daily update's step: state cases found online. What both readers verify goes into found (marked fromNews
    when read from a news report or web search, not the court's own document); what needs a person goes on asks (the
    To check list). Leads are kept in state["stateLeads"]. A news row gives way once the report or a court document
    has the same case. Returns counts."""
    import to_check
    report_rows = list(report["decisions"].values()) + list(report["pending"].values())
    court = [x["case"] for x in report_rows] + [x["case"] for k in ("decisions", "pending") for x in found[k]
                                                if not x.get("fromNews")]
    for k in ("decisions", "pending"):
        found[k] = [x for x in found[k] if not (x.get("fromNews") and known_case(x["case"], court))]
    settled = settle_leads(store, found, court)
    known = court + [x["case"] for k in ("decisions", "pending") for x in found[k] if x.get("fromNews")] + [
        i["case"] for i in store.values() if i.get("case")]  # already asked about on To check
    counts = {"dashboard": 0, "to check": 0, "known": 0, "open": 0, "settled": settled}

    def done(lead, res):
        for k, items in res["adds"].items():
            found.setdefault(k, []).extend(items)
        d, case = res["details"] or {}, data.mask((res["details"] or {}).get("case") or lead["case"])
        if res["what"] == "to check":
            asks.append(to_check.item(
                "statelead:" + "|".join(names(lead["case"])) + "|" + lead["firstSeen"], "State case (found online)",
                res["why"] + " Open the link: if it is a state suit to enforce an IDR award, fill in the Court, State "
                "and Date and answer; if it is not, or it is already on the dashboard, answer Leave out. You can also "
                "save the court's document in State cases\\New.",
                {"case": case, "court": d.get("court") or lead["court"], "level": d.get("level") or "State Trial",
                 "circuit": "", "state": state_code(d.get("state") or lead["state"]), "docket": d.get("docket", ""),
                 "date": d.get("date") or "", "url": res["readFrom"].split(" ")[0] or (lead["sources"] or [""])[0],
                 "payer": data.mask(split_parties(case)[1])},
                "What the search found: " + lead.get("evidence", "")))
        counts["dashboard" if res["adds"] else "to check" if res["what"] == "to check" else
               "known" if res["what"] == "already known" else "open"] += 1
        progress()

    try:
        search(state.setdefault("stateLeads", []), known, reader.system_prompt(report["decisions"].values()), done, say,
               nj_seen=state.setdefault("njOpinionsChecked", []))
    except reader.ReaderError as e:
        say(f"  Claude stopped: {e}. The rest are tried next run.")
    return counts


def settle_leads(store, found, court):
    """Settle the open To check rows of state leads (State case (found online)) whose case the dashboard now has from
    the court's own document or the report (court: those case names), e.g. after a person saved the order in
    State cases\\New. A ruling settles with its result (to_check.dashboard_rows adds no second row for it); otherwise
    Leave out, as the case is already shown. Returns how many."""
    import to_check
    when = dt.date.today().strftime("%b %d, %Y")
    ruled = [x for x in found["decisions"] if not x.get("fromNews") and not x.get("byHand")]
    n = 0
    for key, i in store.items():
        if not key.startswith("statelead:") or i.get("answer"):
            continue
        d = next((x for x in ruled if known_case(i["case"], [x["case"]])), None)
        if d:
            to_check.settle(store, key, d["result"], f"Settled by the daily update on {when}: both readers read the "
                            f"court's ruling of {d.get('date')}.")
        elif known_case(i["case"], court):
            to_check.settle(store, key, to_check.LEAVE, f"Settled by the daily update on {when}: the case is already "
                            "on the dashboard.")
        else:
            continue
        n += 1
    return n


def trial(out=None):
    """A dry run of the daily step: the same search, but results go to State cases\\Search trial\\<date>\\results.xlsx
    and the leads to Search trial\\leads.json. Touches nothing else."""
    base = du.ROOT / state_cases.FOLDER / "Search trial"
    out = Path(out or base / dt.datetime.now().strftime("%Y-%m-%d %H%M"))
    (out / "sources").mkdir(parents=True, exist_ok=True)
    saved_at = base / "leads.json"
    leads = json.loads(saved_at.read_text(encoding="utf-8")) if saved_at.exists() else []
    report = data.load(data.newest_report(du.ROOT / "reports")[0])
    import importlib
    _, payload = importlib.import_module("main").page_data()  # not a plain import: the daily update exe does not need it
    known = [x["case"] for k in ("decisions", "pending") for x in
             (payload[k] if isinstance(payload[k], list) else payload[k].values())]
    rows = []

    def done(lead, res):
        rows.append((lead, res))
        saved_at.write_text(json.dumps(leads, indent=1), encoding="utf-8")  # after each lead: a crash loses nothing

    search(leads, known, reader.system_prompt(report["decisions"].values()), done, du.say, out / "sources")
    saved_at.write_text(json.dumps(leads, indent=1), encoding="utf-8")
    write_results(out / "results.xlsx", rows)
    du.say(f"\nResults: {out / 'results.xlsx'}")
    return rows


def write_results(path, rows):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    head = ["Found by", "Lead: case", "Lead: court", "Lead: state", "Lead: date", "What would happen", "Why / note",
            "Readers: case", "Readers: court", "Readers: state", "Readers: date", "Readers: result", "Readers: holding",
            "Read from", "Sources that could not be used", "Evidence (from the search)"]
    ws.append(head)
    for lead, res in rows:
        got = (res["adds"].get("decisions") or res["adds"].get("pending") or [{}])[0]
        d = res["details"] or {}
        ws.append([lead["foundBy"], lead["case"], lead["court"], lead["state"], lead["date"], res["what"], res["why"],
                   got.get("case") or data.mask(d.get("case", "")), got.get("court") or d.get("court", ""),
                   got.get("state") or d.get("state", ""), got.get("date") or got.get("filed") or d.get("date", ""),
                   got.get("result", ""), got.get("holding", "") or data.mask(d.get("holding", "")), res["readFrom"],
                   "\n".join(res["tried"]), data.mask(lead.get("evidence", ""))])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLMNOP", (14, 40, 30, 8, 11, 34, 50, 40, 30, 8, 11, 10, 60, 40, 50, 60)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    wb.save(path)


if __name__ == "__main__":
    trial()
