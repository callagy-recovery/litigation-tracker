"""Reads the newest NSA enforcement report and works out everything the dashboard shows.

Every text cell is masked as it is read, so nothing downstream (page, logs, tests) ever sees patient initials.
"""
import csv
import datetime as dt
import json
import logging
import re
import shutil
from pathlib import Path

import openpyxl

import insurers

REPORT_RE = re.compile(r"^NSA_Enforcement_Litigation_Dashboard_Source_(\d{4}-\d{2}-\d{2})(?:\s*\(\d+\))?\.xlsx$", re.I)

DECISIONS_SHEET = "NSA Enforcement Decisions"
DECISION_COLS = [
    "Decision ID", "Case Name", "Provider / Claimant", "Payor / Defendant", "Court", "Court Level", "Circuit",
    "Docket No.", "Decision Date", "Court State", "Origin / Plan State", "Outcome", "Heat Map Code",
    "Decision Type", "NSA Private Right of Action?", "FAA §9 Confirmation?", "Award Confirmed?",
    "Award Amount ($)", "Holding / Disposition Summary", "Precedential Weight", "Decision Status / Finality",
    "Appeal / Subsequent History", "Source URL", "Scope Category", "Heat Map Include (Recommended)",
    "Dashboard Caution / Notes", "Verified Through",
]
PENDING_SHEET = "Pending Cases"
PENDING_COLS = ["Case Name", "Provider / Claimant", "Payor / Defendant", "Court", "Court Level", "Circuit",
                "Court State", "Date Filed", "Docket No.", "Status", "Source URL"]
DOL_SHEET = "DOL Actions"
DOL_COLS = ["Title", "Agency", "Type", "Date Issued", "Effective Date", "Status", "Summary", "Source URL"]
METHODOLOGY_SHEET = "Methodology & Coding"

STATE_NAMES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California", "CO": "Colorado",
    "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia", "FL": "Florida", "GA": "Georgia",
    "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
    "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts", "MI": "Michigan",
    "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York", "NC": "North Carolina",
    "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon", "PA": "Pennsylvania",
    "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota", "TN": "Tennessee", "TX": "Texas",
    "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington", "WV": "West Virginia",
    "WI": "Wisconsin", "WY": "Wyoming",
}

CIRCUIT_STATES = {
    "1": ["ME", "MA", "NH", "RI"], "2": ["CT", "NY", "VT"], "3": ["DE", "NJ", "PA"],
    "4": ["MD", "NC", "SC", "VA", "WV"], "5": ["LA", "MS", "TX"], "6": ["KY", "MI", "OH", "TN"],
    "7": ["IL", "IN", "WI"], "8": ["AR", "IA", "MN", "MO", "NE", "ND", "SD"],
    "9": ["AK", "AZ", "CA", "HI", "ID", "MT", "NV", "OR", "WA"], "10": ["CO", "KS", "NM", "OK", "UT", "WY"],
    "11": ["AL", "FL", "GA"], "DC": ["DC"],
}
CIRCUIT_STATES_BY_STATE = {st: c for c, sts in CIRCUIT_STATES.items() for st in sts}
_ORDINAL_WORDS = {"first": "1", "second": "2", "third": "3", "fourth": "4", "fifth": "5", "sixth": "6",
                  "seventh": "7", "eighth": "8", "ninth": "9", "tenth": "10", "eleventh": "11"}

# Patient initials after a/s/o, o/b/o and similar: "a/s/o X.Y." -> "a/s/o ***". The phrase is matched in any case;
# the initials must be capitals, either dotted ("X.Y.", "X. Y.") or two or three bare capitals ("XY").
_INITIALS = r"(?:[A-Z]\.(?:\s?[A-Z]\.?){0,3}|[A-Z]{2,3}(?![A-Za-z]))"
_MASK_RE = re.compile(
    r"(?i:\b(a/s/o|o/b/o|as subrogee of|as assignee of|on behalf of))(\s+)"
    + _INITIALS + r"(?:\s*(?:,|&|and)\s*" + _INITIALS + r")*")


# The firm's own litigation list: the CSV export of app/tools/callagy_cases.sql, one row per docket, saved in callagy\.
CALLAGY_FIELDS = ["case", "court", "docket", "state", "claim_types", "pad_cases", "first_opened", "status"]
_CLAIM_STATE = re.compile(r"^(FL|GA|TX|NJ|NY)")       # claim types such as FLCIOON, GACIOON, TXLIT carry the state
_NJ_DOCKET = re.compile(r"^[A-Z]{3}-+(L|DC)-", re.I)  # NJ Superior Court dockets, including "MRS--L-1750-24"


def load_callagy(*folders):
    """Callagy cases per state from the newest CSV in the folders, or None if there is none. Counts only: the page never
    gets case names. A blank state is taken from the claim type or an NJ docket prefix; the rest are listed as unplaced."""
    files = sorted((f for folder in map(Path, folders) if folder.is_dir() for f in folder.glob("*.csv")),
                   key=lambda f: f.stat().st_mtime)
    if not files:
        return None
    by_state, by_insurer, unplaced = {}, {}, []
    with files[-1].open(encoding="utf-8-sig", errors="replace", newline="") as f:
        for row in csv.reader(f):
            if not any(c.strip() for c in row) or row[0].strip().casefold() == "case_name_constructed":
                continue  # blank line, or a header row if the export has one
            d = dict(zip(CALLAGY_FIELDS, (c.strip() for c in row)))
            st = d.get("state", "").upper()
            if st not in STATE_NAMES:
                m = _CLAIM_STATE.match(d.get("claim_types", "").upper())
                st = m[1] if m else "NJ" if _NJ_DOCKET.match(d.get("docket", "")) else ""
            if st in STATE_NAMES:
                by_state[st] = by_state.get(st, 0) + 1
                for c in insurers.companies(d.get("case", "").partition(" v. ")[2]):  # so the insurer filter applies here too
                    by_insurer.setdefault(c, {})[st] = by_insurer.get(c, {}).get(st, 0) + 1
            else:
                unplaced.append(d.get("docket") or "(no docket)")
    return {"file": files[-1].name, "byState": by_state, "byInsurer": by_insurer, "total": sum(by_state.values()) + len(unplaced),
            "unplaced": unplaced}


class DataFileError(Exception):
    """A data file exists but cannot be read, and its backup cannot either: stop rather than start again from nothing
    and save over it."""


def read_json(path, default):
    """A JSON data file of the daily update. Missing: `default` (a first run). Readable: its contents, and a copy kept
    as <name>.bak, the last good version. Unreadable (half written, a network error): the .bak copy, or DataFileError
    if that cannot be read either - never `default`, which the run would then save over the real data."""
    path = Path(path)
    bak = path.with_name(path.name + ".bak")
    if not path.exists() and not bak.exists():
        return default
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        try:
            obj = json.loads(bak.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            raise DataFileError(f"{path.name} could not be read ({e}) and there is no usable backup ({bak.name}). "
                                "Nothing was changed. Check the network drive, or restore the file, and run again.") from None
        logging.getLogger("lt").warning("%s could not be read (%s); using the backup %s", path.name, e, bak.name)
        return obj
    try:
        shutil.copyfile(path, bak)
    except OSError:
        pass  # a missing backup is not a reason to stop
    return obj


class ReportError(Exception):
    """The report cannot be used; the message names what is wrong and is shown to the user."""


def mask(text):
    return _MASK_RE.sub(lambda m: m[1] + m[2] + "***", text) if isinstance(text, str) else text


def newest_report(folder):
    """(path, date) of the report with the latest date in its name, newest modified time breaking ties; or None."""
    best = None
    for p in Path(folder).iterdir():
        m = REPORT_RE.match(p.name)
        if not m or not p.is_file():
            continue
        try:
            key = (dt.date.fromisoformat(m[1]), p.stat().st_mtime)
        except (ValueError, OSError):
            continue
        if best is None or key > best[0]:
            best = (key, p)
    return (best[1], best[0][0]) if best else None


def norm_docket(text):
    """Every docket number in the text, normalised: '2:22-cv-06614-KM', '4:2022cv06614' -> '22cv6614'; '24-10561' -> '24-10561'.
    '25 Civ. 499 / 1255' -> ['25cv499', '25cv1255']."""
    s = str(text or "").lower()
    out = []
    m = re.search(r"(\d{4}|\d{2})\s*[-\s]?\s*(?:cv|civ\.?|ap|mc)\s*[-\s.]?\s*0*(\d+)", s)
    if m:
        yy = m[1][-2:]
        out.append(f"{yy}cv{int(m[2])}")
        for extra in re.findall(r"/\s*0*(\d+)", s[m.end():]):  # consolidated numbers after the first
            out.append(f"{yy}cv{int(extra)}")
        return out
    m = re.fullmatch(r"\s*(\d{2})-0*(\d+)\s*", s)
    if m:
        return [f"{m[1]}-{int(m[2])}"]
    digits = re.sub(r"\D", "", s)  # state formats vary ("Index No. 604606/2026", "C.A. No. 2025-0011-DH")
    return [digits] if len(digits) >= 4 else ([re.sub(r"[^a-z0-9]", "", s)] if s.strip() else [])


def norm_circuit(value):
    """'3d', '3rd', 'Third', 'Third Circuit', '3' -> '3'; 'D.C. Cir.' -> 'DC'; anything else -> None."""
    s = re.sub(r"\b(?:u\.?s\.?|court of appeals|circuit|cir)\b\.?", " ", str(value or "").lower())
    s = re.sub(r"[\s.()]", "", s)
    if s in ("dc", "districtofcolumbia"):
        return "DC"
    if s in _ORDINAL_WORDS:
        return _ORDINAL_WORDS[s]
    m = re.fullmatch(r"(\d{1,2})(?:st|nd|rd|th|d)?", s)
    return m[1].lstrip("0") if m and 1 <= int(m[1]) <= 11 else None


def circuit_label(key):
    if key == "DC":
        return "D.C. Circuit"
    n = int(key)
    return f"{n}{'th' if 10 <= n % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')} Circuit"


def band(share):
    """Map colour for a win share (0..1). Bands: red <20%, orange <50%, light green <80%, green 80%+."""
    if share < 0.2:
        return "red"
    if share < 0.5:
        return "orange"
    if share < 0.8:
        return "lightgreen"
    return "green"


def state_colour(wins, decisions, pending=0):
    if decisions == 0:
        return "blue" if pending else "grey"
    return band(wins / decisions)


def hatched(decisions):
    """Limited data: a coloured state resting on only 1 or 2 decisions."""
    return decisions in (1, 2)


def _clean(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str):
        return mask(v.strip())
    return v


def _text(v):
    return "" if v is None else str(v).strip()


def _safe_url(v):
    u = _text(v)
    return u if re.match(r"https?://", u, re.I) else ""


def _date(v):
    """ISO date string from a cleaned cell, or '' when it is not a date."""
    s = _text(v)
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return ""


def _one_line(text, limit=180):
    t = " ".join(_text(text).split())
    return t if len(t) <= limit else t[:limit].rsplit(" ", 1)[0].rstrip(",;:") + "…"


def _sheet(wb, name):
    return next((ws for ws in wb.worksheets if ws.title.strip().casefold() == name.casefold()), None)


def _read(ws, cols):
    """Rows as dicts keyed by the expected column names; raises ReportError naming any missing columns."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    idx = {str(h).strip().casefold(): i for i, h in enumerate(header) if h is not None}
    missing = [c for c in cols if c.casefold() not in idx]
    if missing:
        raise ReportError(f'Sheet "{ws.title}" is missing column(s): ' + ", ".join(missing))
    out = []
    for r in rows:
        if not r or all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({c: _clean(r[idx[c.casefold()]]) if idx[c.casefold()] < len(r) else None for c in cols})
    return out


def _methodology(ws):
    items = []
    for r in ws.iter_rows(values_only=True):
        cells = [_text(_clean(v)) for v in r if _text(v)]
        if len(cells) >= 2 and cells[:2] != ["Field / Rule", "Definition"]:
            items.append({"field": cells[0], "text": " ".join(cells[1:])})
        elif len(cells) == 1 and items:  # a lone cell after the table: carry it as a paragraph
            items.append({"field": "", "text": cells[0]})
    return items


def _level(value):
    s = _text(value).casefold()
    return "state" if s.startswith("state") else "federal" if s.startswith("federal") else None


def _view(decisions, pending, level, appellate=False):
    """Per-state figures for one map view. Appellate: each decision colours every state in its circuit."""
    def spread(d):
        if appellate:
            return CIRCUIT_STATES.get(d["circuit"] or "", [])
        return [d["state"]] if d["state"] in STATE_NAMES else []

    ds = [d for d in decisions if d["level"] == level and d["included"] and d["counted"]
          and (not appellate or d["appellate"])]
    ps = [p for p in pending if p["level"] == level and (not appellate or p["appellate"])]
    states = {}
    for ab in STATE_NAMES:
        mine = [d for d in ds if ab in spread(d)]
        mine.sort(key=lambda d: d["date"], reverse=True)
        pend = [p["id"] for p in ps if ab in spread(p)]
        wins = sum(d["win"] for d in mine)
        n = len(mine)
        states[ab] = {
            "name": STATE_NAMES[ab], "decisions": n, "wins": wins, "losses": n - wins,
            "share": wins / n if n else None, "colour": state_colour(wins, n, len(pend)),
            "hatched": hatched(n), "pending": len(pend), "total": n + len(pend), "ids": [d["id"] for d in mine], "pendingIds": pend,
            "via": circuit_label(CIRCUIT_STATES_BY_STATE[ab]) if appellate and ab in CIRCUIT_STATES_BY_STATE else "",
        }
    wins = sum(d["win"] for d in ds)
    excluded = [d for d in decisions if d["level"] == level and not d["included"] and (not appellate or d["appellate"])]
    return {
        "summary": {"included": len(ds), "wins": wins, "losses": len(ds) - wins,
                    "statesWithDecisions": sum(1 for s in states.values() if s["decisions"]),
                    "pending": len(ps), "excluded": len(excluded)},
        "states": states,
        # Every case in this view, newest first, for the "All cases in this view" list.
        "caseIds": [d["id"] for d in sorted(ds, key=lambda d: d["date"], reverse=True)],
        "pendingIds": [p["id"] for p in sorted(ps, key=lambda p: p["filed"] or "", reverse=True)],
    }


AUTO_FILE = "found.json"  # written by the daily update (app/daily_update.py) in auto\
AUTO_STALE_DAYS = 3
AUTO_CAUTION = "Read automatically by the daily update: two AI readers agreed on it. Not reviewed by a person."
HAND_CAUTION = "Checked by a person on a To check list (the AI readers could not confirm it). Not in the report yet."


def _same_case(keys, circuit, state, row):
    """Same docket number in the same circuit or state."""
    return bool(keys & set(norm_docket(row.get("docket")))) and (
        (circuit and circuit == row.get("circuit")) or (state and state == row.get("state")))


def merge_auto(folder, decisions, pending, dol, notes):
    """Add what the daily update found (auto\\found.json) to the report's rows, in place, and return when the daily
    update last ran (ISO date and time, or None). The report always wins: an
    automatic ruling the report already has (same docket, within 3 days) is dropped, and one for a case the report lists
    as pending replaces that pending row. Every text is masked again as it is read."""
    p = Path(folder) / AUTO_FILE
    if not p.is_file():
        return None
    try:
        found = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        notes.append({"kind": "Automatic update not used", "text": f"{p.name} could not be read ({e})."})
        return None
    m = lambda v: mask(_text(v))  # noqa: E731
    count = {"rulings": 0, "pending": 0, "ended": 0, "dol": 0}

    def place(x):
        return set(norm_docket(x.get("docket"))), norm_circuit(x.get("circuit")), _text(x.get("state")).upper()

    def days_apart(a, b):
        try:
            return abs((dt.date.fromisoformat(a) - dt.date.fromisoformat(b)).days)
        except (TypeError, ValueError):
            return 999

    for i, x in enumerate(found.get("decisions") or [], start=1):
        keys, circ, st = place(x)
        if x.get("result") not in ("Win", "Loss") or any(
                _same_case(keys, circ, st, d) and days_apart(d["date"], _date(x.get("date"))) <= 3 for d in decisions):
            # A companion docket decided the same day is its own ruling, as the editor codes it (NSA-042 and NSA-048).
            continue
        pending[:] = [q for q in pending if not _same_case(keys, circ, st, q)]  # decided now
        level, win = _text(x.get("level")), x.get("result") == "Win"
        decisions.append({
            "id": f"AUTO-{i}", "case": m(x.get("case")), "court": m(x.get("court")), "levelText": level,
            "level": _level(level), "appellate": level.casefold() == "federal appellate", "circuit": circ, "state": st,
            "date": _date(x.get("date")), "outcome": "Positive" if win else "Negative", "win": win, "counted": True,
            "status": "Checked by a person" if x.get("byHand") else ("Final" if x.get("final") else "Not final") + " (read automatically)",
            "holding": _one_line(m(x.get("holding")), 400), "url": _safe_url(x.get("url")), "included": True,
            "caution": HAND_CAUTION if x.get("byHand") else AUTO_CAUTION, "payer": m(x.get("payer")),
            "insurers": insurers.companies(m(x.get("payer"))), "docket": _text(x.get("docket")), "auto": True,
            "byHand": bool(x.get("byHand"))})
        count["rulings"] += 1
    for x in found.get("ended") or []:  # settled or voluntarily dismissed: no ruling, no longer pending
        keys, circ, st = place(x)
        before = len(pending)
        pending[:] = [q for q in pending if not _same_case(keys, circ, st, q)]
        count["ended"] += before - len(pending)
    for i, x in enumerate(found.get("pending") or [], start=1):
        keys, circ, st = place(x)
        if any(_same_case(keys, circ, st, r) for r in decisions + pending):
            continue
        level = _text(x.get("level"))
        pending.append({"id": f"AP{i}", "case": m(x.get("case")), "court": m(x.get("court")), "level": _level(level),
                        "appellate": level.casefold() == "federal appellate", "circuit": circ, "state": st,
                        "filed": _date(x.get("filed")), "docket": _text(x.get("docket")), "status": m(x.get("status")),
                        "url": _safe_url(x.get("url")), "payer": m(x.get("payer")),
                        "insurers": insurers.companies(m(x.get("payer"))), "auto": True, "byHand": bool(x.get("byHand"))})
        count["pending"] += 1
    if dol is not None:
        titles = {re.sub(r"[^a-z0-9]", "", a["title"].lower()) for a in dol}
        for x in found.get("dol") or []:
            t = re.sub(r"[^a-z0-9]", "", _text(x.get("title")).lower())
            if not t or t in titles:
                continue
            titles.add(t)
            dol.append({"title": m(x.get("title")), "agency": m(x.get("agency")), "type": m(x.get("type")),
                        "issued": _date(x.get("issued")), "effective": _date(x.get("effective")),
                        "status": m(x.get("status")), "summary": m(x.get("summary")), "url": _safe_url(x.get("url")),
                        "auto": True})
            count["dol"] += 1
        dol.sort(key=lambda a: a["issued"], reverse=True)
    last = _date(found.get("lastRun"))
    notes.append({"kind": "Automatic update", "text": (
        f"Last run {last or 'unknown'}. Added {count['rulings']} rulings, {count['pending']} pending cases and "
        f"{count['dol']} DOL actions; {count['ended']} pending cases ended without a ruling. Rows marked "
        "\"Read automatically\" come from it; the report always wins where both have a case.")})
    if not last or (dt.date.today() - dt.date.fromisoformat(last)).days > AUTO_STALE_DAYS:
        notes.append({"kind": "Automatic updates have stopped", "text": f"The daily update last ran on {last or 'an unknown date'}. "
                      "Run Litigation Tracker Daily Update.exe in the Daily Update folder."})
    run = _text(found.get("lastRun"))
    return run if re.fullmatch(r"\d{4}-\d{2}-\d{2}(T\d{2}:\d{2}(:\d{2})?)?", run) else None


SUMMARY_FILE = "summaries.json"  # two-reader AI summaries written by the daily update (reader.summarise)


def ruling_key(d):
    """One key per ruling across reports and runs: its first docket number and its decision date."""
    return f"{(norm_docket(d.get('docket')) or [''])[0]}|{_date(d.get('date'))}"


def merge_summaries(folder, decisions):
    """Attach each ruling's AI summary, but only while the outcome it was written for is still the ruling's outcome."""
    p = Path(folder) / SUMMARY_FILE
    try:
        found = json.loads(p.read_text(encoding="utf-8")) if p.is_file() else {}
    except (OSError, ValueError):
        return
    for d in decisions:
        s = found.get(ruling_key(d))
        if s and s.get("summary") and s.get("outcome") == ("Win" if d.get("win") else "Loss") and d.get("counted"):
            d["summary"] = mask(_text(s.get("summary")))


def views(decisions, pending, only=None):
    """The three map views. With only (a list of standard insurer names), just the cases naming one of them."""
    if only:
        only = set(only)
        decisions = [d for d in decisions if only & set(d["insurers"])]
        pending = [p for p in pending if only & set(p["insurers"])]
    return {
        "state": _view(decisions, pending, "state"),
        "federal": _view(decisions, pending, "federal"),
        "appellate": _view(decisions, pending, "federal", appellate=True),
    }



def load(path, auto=None):
    """Everything the page needs from one report, plus the daily update's finds from the auto folder if given.
    Raises ReportError for a missing sheet or column."""
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # corrupt, locked mid-save, not really xlsx
        raise ReportError(f"Could not open {path.name}: {e}") from e
    try:
        ws = _sheet(wb, DECISIONS_SHEET)
        if ws is None:
            raise ReportError(f'Required sheet "{DECISIONS_SHEET}" is missing from {path.name}.')
        rows = _read(ws, DECISION_COLS)
        notes = []
        optional = {}
        for name, cols in ((PENDING_SHEET, PENDING_COLS), (DOL_SHEET, DOL_COLS)):
            s = _sheet(wb, name)
            try:
                optional[name] = _read(s, cols) if s is not None else None
            except ReportError as e:
                notes.append({"kind": "Sheet not used", "text": f"{e} The sheet is ignored until that is fixed."})
                optional[name] = None
        m = _sheet(wb, METHODOLOGY_SHEET)
        methodology = _methodology(m) if m is not None else None
    finally:
        wb.close()

    decisions = []
    seen = set()
    for i, r in enumerate(rows, start=2):
        rid = _text(r["Decision ID"]) or f"row {i}"
        if rid in seen:
            notes.append({"kind": "Duplicate Decision ID", "text": f"{rid} appears more than once; the copy on row {i} is shown as \"{rid} (row {i})\"."})
            rid = f"{rid} (row {i})"
        seen.add(rid)
        level = _level(r["Court Level"])
        outcome = _text(r["Outcome"]).casefold()
        include_raw = _text(r["Heat Map Include (Recommended)"]).casefold()
        included = include_raw == "yes"
        circuit = norm_circuit(r["Circuit"])
        state = _text(r["Court State"]).upper()
        if level is None:
            notes.append({"kind": "Unknown Court Level", "text": f"{rid}: \"{_text(r['Court Level'])}\" is neither State nor Federal, so it is not mapped."})
        if outcome not in ("positive", "negative"):
            notes.append({"kind": "Unknown Outcome", "text": f"{rid}: \"{_text(r['Outcome'])}\" is not Positive or Negative, so it is not counted."})
        if include_raw not in ("yes", "no"):
            notes.append({"kind": "Unknown Heat Map Include", "text": f"{rid}: \"{_text(r['Heat Map Include (Recommended)'])}\" is treated as No."})
        if level == "federal" and circuit is None:
            notes.append({"kind": "Unknown Circuit", "text": f"{rid}: \"{_text(r['Circuit'])}\" is not a recognised circuit."})
        if included and level and state not in STATE_NAMES:
            notes.append({"kind": "Unknown Court State", "text": f"{rid}: \"{state}\" is not one of the 50 states or DC, so it is not mapped."})
        code = r["Heat Map Code"]
        expected = {"positive": 1, "negative": -1}.get(outcome)
        try:
            agrees = expected is not None and float(code) == expected
        except (TypeError, ValueError):
            agrees = False
        if expected is not None and not agrees:
            notes.append({"kind": "Outcome and Heat Map Code disagree",
                          "text": f"{rid}: Outcome is {_text(r['Outcome'])} but Heat Map Code is {_text(code) or 'blank'}. Outcome is used."})
        decisions.append({
            "id": rid, "case": _text(r["Case Name"]), "court": _text(r["Court"]), "levelText": _text(r["Court Level"]),
            "level": level, "appellate": _text(r["Court Level"]).casefold() == "federal appellate",
            "circuit": circuit, "state": state, "date": _date(r["Decision Date"]),
            "outcome": "Positive" if outcome == "positive" else "Negative" if outcome == "negative" else _text(r["Outcome"]),
            "win": outcome == "positive", "counted": outcome in ("positive", "negative"),
            "status": _text(r["Decision Status / Finality"]), "holding": _one_line(r["Holding / Disposition Summary"]),
            "url": _safe_url(r["Source URL"]), "included": included, "caution": _text(r["Dashboard Caution / Notes"]),
            "payer": _text(r["Payor / Defendant"]), "insurers": insurers.companies(_text(r["Payor / Defendant"])),
            "docket": _text(r["Docket No."]),
        })

    pending = []
    for i, r in enumerate(optional[PENDING_SHEET] or [], start=2):
        pending.append({
            "id": f"P{i}", "case": _text(r["Case Name"]), "court": _text(r["Court"]), "level": _level(r["Court Level"]),
            "appellate": _text(r["Court Level"]).casefold() == "federal appellate", "circuit": norm_circuit(r["Circuit"]),
            "state": _text(r["Court State"]).upper(), "filed": _date(r["Date Filed"]), "docket": _text(r["Docket No."]),
            "status": _text(r["Status"]), "url": _safe_url(r["Source URL"]),
            "payer": _text(r["Payor / Defendant"]), "insurers": insurers.companies(_text(r["Payor / Defendant"])),
        })
    dol = None
    if optional[DOL_SHEET] is not None:
        dol = [{"title": _text(r["Title"]), "agency": _text(r["Agency"]), "type": _text(r["Type"]),
                "issued": _date(r["Date Issued"]) or _text(r["Date Issued"]),
                "effective": _date(r["Effective Date"]) or _text(r["Effective Date"]),
                "status": _text(r["Status"]), "summary": _text(r["Summary"]), "url": _safe_url(r["Source URL"])}
               for r in optional[DOL_SHEET]]
        dol.sort(key=lambda a: a["issued"], reverse=True)
    auto_run = None
    if auto:
        auto_run = merge_auto(auto, decisions, pending, dol, notes)
        merge_summaries(auto, decisions)

    included = [d for d in decisions if d["included"]]
    return {
        "counts": {"read": len(decisions), "included": len(included), "excluded": len(decisions) - len(included),
                   "wins": sum(d["win"] for d in included if d["counted"]),
                   "losses": sum(not d["win"] for d in included if d["counted"])},
        "views": views(decisions, pending),
        "decisions": {d["id"]: d for d in decisions},
        "excludedList": [d for d in decisions if not d["included"]],
        "pending": {p["id"]: p for p in pending},
        "hasPendingSheet": optional[PENDING_SHEET] is not None,
        "dol": dol,
        "methodology": methodology,
        "notes": notes,
        "autoRun": auto_run,  # when the daily update last added to this data (ISO), or None
    }
