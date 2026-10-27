"""The "To check" list: what the two AI readers could not confirm is added at the bottom of one file,
"To check\\To check.xlsx", for a person (or the legal team) to answer; "Added on" is the day it was found. The next
daily update reads the answers and shows them on the dashboard, marked "Checked by a person". A row is taken off the
To check sheet once it is settled - answered by a person, or decided by the daily update itself (`settle`) - and kept
on the file's "Settled" sheet, where its answer can still be changed. (Before 2026-10-27 each day had its own file;
those are in "To check\\Done\\". Any other .xlsx left in To check\\ is still read and tidied.)

Each item's full details stay in auto\\to_check.json (the workbook only shows them), keyed by a stable key such as
"docket:<CourtListener id>", "rd:<entry id>" or "state:<file name>", so nothing is asked twice.
"""
import datetime as dt
import json
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import data

FOLDER, DONE, STORE, FILE = "To check", "Done", "to_check.json", "To check.xlsx"
PENDING, WIN, LOSS, LEAVE = "Pending case", "Win", "Loss", "Leave out"
ANSWERS = (PENDING, WIN, LOSS, LEAVE)
COLUMNS = ["Case", "Court", "Docket number", "State", "Date", "What it is", "Why it needs a person",
           "What the readers found", "Link", "Your answer", "Note (optional)", "Key", "Added on"]
SETTLED, SETTLED_COLUMNS = "Settled", ["Settled on", "Settled by"]
WIDTHS = (46, 32, 16, 7, 12, 22, 44, 50, 44, 16, 36, 4, 12)
ANSWER_COL, KEY_COL = COLUMNS.index("Your answer") + 1, COLUMNS.index("Key") + 1  # 1-based, for openpyxl
HEAD = PatternFill("solid", fgColor="1B2F6B")
ANSWER = PatternFill("solid", fgColor="FFF7E6")
READ_ME = [
    "What this is: cases and rulings the two AI readers could not confirm on their own. Everything they did confirm is "
    "already on the dashboard.",
    "What to do: for each row, open the link, then pick Your answer: Pending case (a suit to enforce an IDR award, not "
    "decided yet), Win or Loss (a ruling on enforcing the award), or Leave out (not one for the tracker). Add a note if "
    "you like; for a Win or Loss it is shown as the holding. If the Case, Court, Docket number, State or Date is wrong, "
    "correct it in the row. Before answering from a saved page (Justia, CourtListener), check its docket number "
    "matches this row: one plaintiff often has several cases against the same insurer.",
    "Then: save the file here, under the same name, and close it. The next daily update puts the answers on the "
    "dashboard, marked 'Checked by a person', and moves the answered rows to the Settled sheet. Leave a row empty if "
    "you are not sure; it stays on the list. The daily update also settles rows by itself when it finds the answer (a "
    "ruling both readers agree on, a settlement, or a case that is not an enforcement suit); those move to the Settled "
    "sheet too, with the reason in the Note. To change any answer, change it on the Settled sheet.",
    "New rows: each daily update adds what it could not confirm at the bottom of the To check sheet; 'Added on' is the "
    "day it was found. If this file is open in Excel during a run, its new rows are added on the next run instead.",
]


def folder(root):
    return Path(root) / FOLDER


def load(root):
    return data.read_json(Path(root) / "auto" / STORE, {})  # unreadable: the backup, or the run stops


def save(root, store):
    p = Path(root) / "auto" / STORE
    tmp = p.with_suffix(".tmp")
    tmp.write_text(json.dumps(store, indent=1, ensure_ascii=False), encoding="utf-8")
    data.swap_in(tmp, p)


def item(key, kind, why, row, readers=""):
    """One thing for a person to check. row has the dashboard fields: case, court, level, circuit, state, docket, date,
    url, payer (and docketId for a CourtListener docket)."""
    return {"key": key, "kind": kind, "why": data.mask(why), "readers": data.mask(readers),
            **{k: (data.mask(v) if k in ("case", "payer") else v) for k, v in row.items()}}


def add(root, store, items, today=None):
    """Record the items not asked before, then add every recorded row not yet in the file at the bottom of
    To check.xlsx. Returns (file or None, rows added). If the file is open in Excel nothing is lost: the rows stay
    recorded as unwritten and are added by the next call (the next run)."""
    today = today or dt.date.today()
    for i in items:
        if i["key"] not in store:
            store[i["key"]] = {**i, "asked": today.isoformat(), "unwritten": True}
    waiting = [i for i in store.values() if i.get("unwritten") and not i.get("answer")]
    if not waiting:
        return None, 0
    folder(root).mkdir(exist_ok=True)
    f = folder(root) / FILE
    try:
        _write(f, waiting)
    except PermissionError:  # open in Excel
        return None, 0
    for i in waiting:
        i.pop("unwritten")
        i["file"] = FILE
    return f, len(waiting)


def unwritten(store):
    """Rows waiting to go into To check.xlsx because it was open in Excel."""
    return sum(1 for i in store.values() if i.get("unwritten") and not i.get("answer"))


def _write(f, new):
    """Append the rows at the bottom of f's To check sheet (making the file the first time)."""
    if f.exists():
        wb = openpyxl.load_workbook(f)
        ws = wb["To check"]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "To check"
        ws.append(COLUMNS)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.column_dimensions[_letter(KEY_COL)].hidden = True  # the key, so answers can be matched even if rows are sorted
        ws.freeze_panes = "A2"
        dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{_letter(ANSWER_COL)}2:{_letter(ANSWER_COL)}2000")
        how = wb.create_sheet("Read me")
        how.column_dimensions["A"].width = 120
        for line in READ_ME:
            how.append([line])
            how.cell(how.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    head = [c.value for c in ws[1]]  # written by the file's own headings, so an older layout still lines up
    for i in new:
        values = {"Case": i["case"], "Court": i.get("court"), "Docket number": i.get("docket"), "State": i.get("state"),
                  "Date": i.get("date"), "What it is": i["kind"], "Why it needs a person": i["why"],
                  "What the readers found": i["readers"], "Link": i.get("url"), "Key": i["key"], "Added on": i.get("asked")}
        ws.append([values.get(h) for h in head])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if "Your answer" in head:
            ws.cell(ws.max_row, head.index("Your answer") + 1).fill = ANSWER
    wb.save(f)


def _letter(n):
    return openpyxl.utils.get_column_letter(n)


def collect(root, store, say=lambda m: None):
    """Read Your answer from every file in To check\\ (both its To check and Settled sheets) into the store. Returns
    the keys whose answer is new or changed. An answer a person changes overrides one the daily update gave."""
    changed = []
    for f in sorted(folder(root).glob("*.xlsx")):
        if f.name.startswith("~$"):  # Excel's lock file
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            sheets = [list(wb[n].iter_rows(values_only=True)) for n in ("To check", SETTLED) if n in wb.sheetnames]
            wb.close()
        except Exception as e:  # damaged, or not one of ours
            say(f"  {f.name} could not be read ({e}); left as it is.")
            continue
        for rows in sheets:
            changed += _answers(rows, store)
    return changed


def _answers(rows, store):
    if not rows:
        return []
    col = {h: n for n, h in enumerate(rows[0]) if h}
    if "Key" not in col or "Your answer" not in col:
        return []
    changed = []
    for r in rows[1:]:
        key = r[col["Key"]] if len(r) > col["Key"] else None
        if not key or key not in store:
            continue
        ans = str(r[col["Your answer"]] or "").strip()
        ans = next((a for a in ANSWERS if a.casefold() == ans.casefold()), "")
        if not ans:
            continue
        note = str(r[col["Note (optional)"]] or "").strip() if "Note (optional)" in col else ""
        # The person may have corrected what the readers made of the case.
        fixes = {k: _cell(r[col[h]]) for h, k in (("Case", "case"), ("Court", "court"), ("Docket number", "docket"),
                                                  ("State", "state"), ("Date", "date")) if h in col and _cell(r[col[h]])}
        fixes = {k: (data.mask(v) if k == "case" else v.upper() if k == "state" else v) for k, v in fixes.items()}
        if store[key].get("answer") != ans or store[key].get("note", "") != data.mask(note) or any(store[key].get(k) != v for k, v in fixes.items()):
            store[key].update(answer=ans, note=data.mask(note), answered=dt.date.today().isoformat(), **fixes)
            store[key].pop("auto", None)  # a person's answer from now on
            changed.append(key)
    return changed


def settle(store, key, answer, note):
    """The daily update found the answer itself (a ruling both readers agree on, a settlement, not an enforcement
    suit): the row leaves the To check sheet at the end of the run. What the readers found is already on the dashboard,
    so dashboard_rows adds nothing for it."""
    store[key].update(answer=answer, note=data.mask(note), answered=dt.date.today().isoformat(), auto=True)


def tidy(root, store, say=lambda m: None):
    """Move every settled row from the To check sheet to the Settled sheet of its file; an older file (not
    To check.xlsx, which always stays) with no rows left to check moves to To check\\Done\\. A file open in Excel is
    left as it is and tidied next time. Returns rows moved."""
    moved = 0
    for f in sorted(folder(root).glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(f)
        except Exception as e:
            say(f"  {f.name} could not be read ({e}); left as it is.")
            continue
        if "To check" not in wb.sheetnames:
            continue
        ws = wb["To check"]
        col = {c.value: c.column for c in ws[1] if c.value}
        if "Key" not in col:
            continue
        done = wb[SETTLED] if SETTLED in wb.sheetnames else _settled_sheet(wb)
        gone, left = [], 0
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, col["Key"]).value
            if not key:
                continue
            i = store.get(key)
            if not i or not i.get("answer"):
                left += 1
                continue
            values = [ws.cell(r, col[h]).value if h in col else None for h in COLUMNS]
            values[COLUMNS.index("Your answer")], values[COLUMNS.index("Note (optional)")] = i["answer"], i.get("note") or None
            done.append(values + [i.get("answered"), "the daily update" if i.get("auto") else "a person"])
            for c in done[done.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            gone.append(r)
        if not gone and left:
            continue
        for r in reversed(gone):
            ws.delete_rows(r)
        try:
            wb.save(f)
        except PermissionError:
            say(f"  {f.name} is open in Excel: its settled rows are moved next time.")
            continue
        moved += len(gone)
        if not left and f.name != FILE:
            (f.parent / DONE).mkdir(exist_ok=True)
            try:
                os.replace(f, f.parent / DONE / f.name)
                say(f"  every row of {f.name} is settled: moved to {FOLDER}\\{DONE}")
            except OSError:
                say(f"  every row of {f.name} is settled; close it in Excel and it moves to {DONE} next time.")
    return moved


def _settled_sheet(wb):
    ws = wb.create_sheet(SETTLED, 1)
    ws.append(COLUMNS + SETTLED_COLUMNS)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(WIDTHS + (12, 16), start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.column_dimensions[_letter(KEY_COL)].hidden = True
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{_letter(ANSWER_COL)}2:{_letter(ANSWER_COL)}5000")
    return ws


def _cell(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat() if isinstance(v, dt.date) else str(v or "").strip()


def dashboard_rows(store):
    """What the answered items add to found.json: ({"decisions": [...], "pending": [...]}, keys answered Leave out)."""
    out, leave = {"decisions": [], "pending": []}, set()
    for key, i in store.items():
        a = i.get("answer")
        if not a:
            continue
        if a == LEAVE:
            leave.add(key)
            continue
        if i.get("auto"):
            continue  # settled by the daily update: the readers' own row is already on the dashboard
        row = {k: i.get(k) for k in ("case", "court", "level", "circuit", "state", "docket", "url", "payer")}
        row.update(byHand=True, toCheck=key)
        if a == PENDING:
            out["pending"].append({**row, "filed": i.get("filed") or i.get("date"), "docketId": i.get("docketId"),
                                   "status": "Checked by a person: an enforcement suit with no ruling yet."})
        else:
            out["decisions"].append({**row, "date": i.get("date"), "result": a, "final": None,
                                     "holding": i.get("note") or f"A {a.lower()} for the provider, checked by a person."})
    return out, leave


def open_count(store):
    return sum(1 for i in store.values() if not i.get("answer"))
