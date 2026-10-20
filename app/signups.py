"""Newsletter sign-ups: one Excel list, newsletter\\Newsletter sign-ups.xlsx, that the dashboard's Sign up button
writes and the newsletter reads.

Columns: Email, Name, Windows user, Signed up, Status (Subscribed / Unsubscribed), Changed. An email is on the list
once (compared without regard to capitals); signing up again only turns an Unsubscribed row back on. Only work
addresses are accepted. Writes take a short lock file, so two people signing up at the same moment cannot overwrite
each other, and are saved to a temporary file first, so a half-written list is never left behind. If the list is open
in Excel, Windows will not let it be replaced; the caller is told to try again.
"""
import datetime as dt
import os
import re
import time
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

FOLDER, FILE = "newsletter", "Newsletter sign-ups.xlsx"
HEADERS = ["Email", "Name", "Windows user", "Signed up", "Status", "Changed"]
DOMAINS = ("callagyrecovery.com", "callagylaw.com")
_EMAIL = re.compile(r"^[a-z0-9._%+-]+@([a-z0-9-]+\.)+[a-z]{2,}$")


class SignupError(Exception):
    """Shown to the person signing up."""


def path(root):
    return Path(root) / FOLDER / FILE


def clean_email(email):
    e = (email or "").strip().lower()
    if not _EMAIL.match(e):
        raise SignupError("That doesn't look like an email address.")
    if e.rsplit("@", 1)[1] not in DOMAINS:
        raise SignupError("Please use your work address (" + " or ".join("@" + d for d in DOMAINS) + ").")
    return e


def guess(user):
    """Email and name from the Windows user name: 'jane.doe' -> ('jane.doe@callagyrecovery.com', 'Jane Doe')."""
    u = (user or "").strip().lower()
    name = " ".join(w.capitalize() for w in re.split(r"[._-]+", u) if w)
    return (f"{u}@{DOMAINS[0]}" if re.fullmatch(r"[a-z0-9._-]+", u) else ""), name


def _rows(p):
    if not p.is_file():
        return []
    wb = openpyxl.load_workbook(p, read_only=True)
    try:
        rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    return [dict(zip(HEADERS, r)) for r in rows if r and r[0]]


def status(root, email):
    """'Subscribed', 'Unsubscribed' or None for this address."""
    e = (email or "").strip().lower()
    return next((r["Status"] for r in _rows(path(root)) if str(r["Email"]).strip().lower() == e), None)


def recipients(root):
    """Everyone subscribed, in sign-up order."""
    return [str(r["Email"]).strip() for r in _rows(path(root)) if r["Status"] == "Subscribed"]


class _Lock:
    def __init__(self, p):
        self.lock = p.with_name(p.name + ".lock")

    def __enter__(self):
        for _ in range(50):  # about 10 seconds
            try:
                os.close(os.open(self.lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY))
                return self
            except FileExistsError:
                try:
                    if time.time() - self.lock.stat().st_mtime > 60:  # left behind by a crash
                        self.lock.unlink()
                        continue
                except OSError:
                    pass
                time.sleep(0.2)
        raise SignupError("Someone else is signing up right now; please try again.")

    def __exit__(self, *exc):
        try:
            self.lock.unlink()
        except OSError:
            pass


def _save(p, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sign-ups"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="276BA8")
        c.alignment = Alignment(vertical="top")
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    for col, w in zip("ABCDEF", (36, 22, 18, 12, 14, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    tmp = p.with_name("~" + p.stem + ".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, p)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise SignupError("The sign-up list is open in Excel right now; please try again in a minute.") from None


def subscribe(root, email, name="", user="", today=None):
    """'added', 'resubscribed' or 'already'."""
    e, p, today = clean_email(email), path(root), (today or dt.date.today()).isoformat()
    p.parent.mkdir(exist_ok=True)
    with _Lock(p):
        rows = _rows(p)
        row = next((r for r in rows if str(r["Email"]).strip().lower() == e), None)
        if row and row["Status"] == "Subscribed":
            return "already"
        if row:
            row.update({"Status": "Subscribed", "Changed": today})
            _save(p, rows)
            return "resubscribed"
        rows.append({"Email": e, "Name": name, "Windows user": user, "Signed up": today, "Status": "Subscribed", "Changed": today})
        _save(p, rows)
        return "added"


def unsubscribe(root, email, today=None):
    """'removed' or 'not on the list'."""
    e, p, today = clean_email(email), path(root), (today or dt.date.today()).isoformat()
    with _Lock(p):
        rows = _rows(p)
        row = next((r for r in rows if str(r["Email"]).strip().lower() == e and r["Status"] == "Subscribed"), None)
        if not row:
            return "not on the list"
        row.update({"Status": "Unsubscribed", "Changed": today})
        _save(p, rows)
        return "removed"
