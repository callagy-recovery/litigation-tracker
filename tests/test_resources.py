import logging
import os
import time

import docx
from pypdf import PdfWriter

import resources


def make_pdf(path, lines, font_matrix=False):
    """Minimal one-page PDF. lines: [(font size, text, y)]. font_matrix=True sets size via Tm scaling instead of Tf."""
    ops = []
    for size, text, y in lines:
        if font_matrix:
            ops.append(f"BT /F1 1 Tf {size} 0 0 {size} 72 {y} Tm ({text}) Tj ET")
        else:
            ops.append(f"BT /F1 {size} Tf 1 0 0 1 72 {y} Tm ({text}) Tj ET")
    stream = "\n".join(ops).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


def test_pdf_title_property(tmp_path):
    w = PdfWriter()
    w.add_blank_page(612, 792)
    w.add_metadata({"/Title": "  IDR Enforcement   Update  "})
    p = tmp_path / "a.pdf"
    with open(p, "wb") as f:
        w.write(f)
    assert resources.pdf_title(p) == "IDR Enforcement Update"


def test_pdf_falls_back_to_largest_text_on_page_one(tmp_path):
    p = tmp_path / "b.pdf"
    make_pdf(p, [(10, "Small print at the top", 760), (24, "Second Circuit Rules on IDR Awards", 700),
                 (12, "Body text follows", 650)])
    assert resources.pdf_title(p) == "Second Circuit Rules on IDR Awards"


def test_pdf_largest_text_counts_text_matrix_scaling(tmp_path):
    p = tmp_path / "c.pdf"
    make_pdf(p, [(12, "Body", 700), (30, "Scaled Headline", 650)], font_matrix=True)
    assert resources.pdf_title(p) == "Scaled Headline"


def test_pdf_with_no_title_or_text_falls_back_to_file_name(tmp_path):
    w = PdfWriter()
    w.add_blank_page(612, 792)
    p = tmp_path / "Blank Memo.pdf"
    with open(p, "wb") as f:
        w.write(f)
    assert resources.pdf_title(p) == ""
    items, problems = resources.scan(tmp_path)
    assert items[0]["title"] == "Blank Memo" and not problems


def test_docx_core_title(tmp_path):
    d = docx.Document()
    d.core_properties.title = "NSA Litigation Memo"
    d.add_heading("A heading that should lose", level=1)
    p = tmp_path / "a.docx"
    d.save(p)
    assert resources.docx_title(p) == "NSA Litigation Memo"


def test_docx_falls_back_to_first_heading(tmp_path):
    d = docx.Document()
    d.core_properties.title = ""
    d.add_paragraph("Ordinary opening paragraph")
    d.add_heading("Fifth Circuit Guardian Flight Summary", level=2)
    d.add_heading("Later heading", level=1)
    p = tmp_path / "b.docx"
    d.save(p)
    assert resources.docx_title(p) == "Fifth Circuit Guardian Flight Summary"


def test_docx_without_title_or_heading_falls_back_to_file_name(tmp_path):
    d = docx.Document()
    d.core_properties.title = ""
    d.add_paragraph("Just text")
    d.save(tmp_path / "Client Notes.docx")
    items, _ = resources.scan(tmp_path)
    assert items[0]["title"] == "Client Notes" and items[0]["type"] == "Word"


def test_html_title_tag(tmp_path):
    p = tmp_path / "a.html"
    p.write_text("<html><head><title> DOL &amp; IDR  guidance </title></head><body><h1>Heading</h1></body></html>", encoding="utf-8")
    assert resources.html_title(p) == "DOL & IDR guidance"


def test_html_falls_back_to_first_h1(tmp_path):
    p = tmp_path / "b.htm"
    p.write_text("<html><head><title>  </title></head><body><h1>Court <em>confirms</em> award</h1><h1>Second</h1></body></html>", encoding="utf-8")
    assert resources.html_title(p) == "Court confirms award"


def test_html_without_title_or_h1_falls_back_to_file_name(tmp_path):
    (tmp_path / "Saved Article.html").write_bytes(b"<html><body><p>caf\xe9</p></body></html>")  # cp1252, not utf-8
    items, problems = resources.scan(tmp_path)
    assert items[0]["title"] == "Saved Article" and items[0]["type"] == "HTML" and not problems


def test_unreadable_files_are_listed_under_their_name_and_logged(tmp_path, caplog):
    (tmp_path / "Broken Brief.pdf").write_bytes(b"not a pdf at all")
    (tmp_path / "Broken Memo.docx").write_bytes(b"not a zip")
    with caplog.at_level(logging.WARNING, logger="lt"):
        items, problems = resources.scan(tmp_path)
    assert sorted(i["title"] for i in items) == ["Broken Brief", "Broken Memo"]
    assert len(problems) == 2
    assert "Broken Brief.pdf" in caplog.text and "Broken Memo.docx" in caplog.text


def test_scan_lists_supported_types_newest_first_and_skips_the_rest(tmp_path):
    now = time.time()
    for name, age in (("old.html", 300), ("new.htm", 10), ("middle.pdf", 100), ("notes.txt", 0),
                      ("~$lock.docx", 0), ("Thumbs.db", 0)):
        p = tmp_path / name
        if name.endswith(".pdf"):
            make_pdf(p, [(20, "Middle", 700)])
        else:
            p.write_text("<title>t</title>", encoding="utf-8")
        os.utime(p, (now - age, now - age))
    items, _ = resources.scan(tmp_path)
    assert [i["name"] for i in items] == ["new.htm", "middle.pdf", "old.html"]
    assert all(len(i["token"]) == 16 for i in items)


def test_titles_are_masked(tmp_path):
    (tmp_path / "Clinic a-s-o.html").write_text("<title>Clinic a/s/o X.Y. v. Plan</title>", encoding="utf-8")
    items, _ = resources.scan(tmp_path)
    assert items[0]["title"] == "Clinic a/s/o *** v. Plan"


def test_missing_folder_gives_empty_list(tmp_path):
    assert resources.scan(tmp_path / "nope") == ([], [])


def test_links_csv_is_listed_with_files_newest_first(tmp_path):
    (tmp_path / "links.csv").write_text(
        "Title,Type,Source,Date,URL,Summary\n"
        '"Clinic a/s/o X.Y. v. Plan, ruling",Court opinion,D.N.J.,2026-09-01,https://example.org/a.pdf,"Award confirmed."\n'
        "Agency notice,government,CMS,2025-01-02,https://cms.gov/x,\n"
        "Undated news,News,HFMA,,https://hfma.org/y,\n"
        "No link here,News,HFMA,2026-01-01,ftp://bad,\n"
        ",,,,,\n", encoding="utf-8")
    pdf = tmp_path / "note.pdf"
    pdf.write_bytes(b"%PDF-1.4 broken")  # unreadable: still listed, dated today
    items, problems = resources.listing(tmp_path)
    links = [i for i in items if i["type"] == "Link"]
    assert [i["title"] for i in links] == ["Clinic a/s/o *** v. Plan, ruling", "Agency notice", "Undated news"]
    assert links[1]["category"] == "Government"           # type matched ignoring case
    assert items[0]["type"] == "PDF" and items[-1]["title"] == "Undated news"  # today's file first, undated link last
    assert any("row 5 skipped" in p for p in problems)    # the ftp row; the blank row is ignored silently
    assert resources.links(tmp_path / "missing") == ([], [])
