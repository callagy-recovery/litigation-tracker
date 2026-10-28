"""The public website (app/publish.py -> docs\\): nothing private in it, and its in-page insurer filter
(app/static/views.js) gives exactly what the dashboard's own Python (data.views) gives."""
import io
import json
import re
import shutil
import subprocess

import openpyxl
import pytest

import data
import main
from conftest import REPORT, ROOT


def payload_of(page):
    return json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>', page, re.S)[1].replace("<\\/", "</"))


@pytest.mark.skipif(not REPORT.exists(), reason="acceptance report not in reports\\")
def test_the_website_carries_nothing_private(tmp_path, monkeypatch):
    import publish
    reports, callagy = tmp_path / "reports", tmp_path / "callagy"
    reports.mkdir(), callagy.mkdir()
    shutil.copy(REPORT, reports / REPORT.name)
    (callagy / "cases.csv").write_text("Secret Client Medical PC v. Plan,District Court,1:25-cv-1,NJ,NJLIT,1,2025-01-01,Open\n",
                                       encoding="utf-8")
    res = tmp_path / "resources"
    res.mkdir()
    (res / "Internal memo.html").write_text("<title>Internal memo</title>", encoding="utf-8")
    for name, value in (("REPORTS", reports), ("CALLAGY", callagy), ("RESOURCES", res), ("AUTO", tmp_path / "auto")):
        monkeypatch.setattr(main, name, value)
    auto = tmp_path / "auto"
    auto.mkdir()
    (auto / data.AUTO_FILE).write_text('{"rows": [{"case": "Clinic a/s/o X.Y. v. Plan"}]}', encoding="utf-8")
    out = tmp_path / "docs"
    (out / "static").mkdir(parents=True)
    (out / "old-logo.png").write_bytes(b"x")  # left from an earlier publish
    files = publish.publish(out)
    assert {"index.html", "robots.txt", ".nojekyll", "static/app.js", "static/views.js", "static/us_states.json",
            publish.DATA_XLSX, "data/daily-update-results.json"} <= set(files)
    assert "data/ai-summaries.json" not in files  # no summaries file in auto\, so none published
    assert "a/s/o ***" in (out / "data/daily-update-results.json").read_text(encoding="utf-8")
    wb = openpyxl.load_workbook(out / publish.DATA_XLSX)
    texts = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert texts and all(data.mask(t) == t for t in texts if not t.startswith("="))  # nothing left to mask
    assert any("a/s/o ***" in t for t in texts)
    assert wb.properties.creator == wb.properties.lastModifiedBy == "Litigation Tracker"
    assert not any(c.comment for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    assert payload_of((out / "index.html").read_text(encoding="utf-8"))["dataDownload"] == publish.DATA_XLSX
    assert not (out / "old-logo.png").exists()
    page = (out / "index.html").read_text(encoding="utf-8")
    d = payload_of(page)
    assert d["published"] and d["callagy"] is None and d["version"] is None and "token" not in d["signup"]
    assert "Secret Client" not in page and "Internal memo" not in page and "/download/" not in page
    assert 'name="robots" content="noindex' in page and (out / "robots.txt").read_text() == "User-agent: *\nDisallow: /\n"
    assert 'src="static/app.js"' in page and 'href="/static' not in page  # works from a GitHub project address


def test_staff_names_and_internal_paths_are_reworded():
    import publish
    note = {"n": ["Added 2026-10-16 from Katie Varner's review of 2026-10-12.",
                  "Items still under review are in drafts\\Katie_Review_2_2026-10-16.xlsx.", "Clinic a/s/o X.Y. v. Plan"]}
    assert publish.public(note) == {"n": ["Added 2026-10-16 from an internal review of 2026-10-12.",
                                          "Items still under review are in an internal file.", "Clinic a/s/o *** v. Plan"]}


@pytest.mark.skipif(shutil.which("git") is None, reason="git not installed")
def test_push_sends_changes_only_from_main(tmp_path, monkeypatch):
    import publish
    remote, proj = tmp_path / "remote.git", tmp_path / "proj"
    git = lambda *a, cwd=proj: subprocess.run(publish.GIT + list(a), cwd=cwd, capture_output=True, text=True, check=True).stdout  # noqa: E731
    git("init", "-q", "--bare", str(remote), cwd=tmp_path)
    proj.mkdir()
    git("init", "-q", "-b", "main")
    git("remote", "add", "origin", str(remote))
    (proj / ".gitignore").write_text("CLAUDE.md\n", encoding="utf-8")
    (proj / "CLAUDE.md").write_text("private", encoding="utf-8")
    (proj / "notes.txt").write_text("not a public path", encoding="utf-8")
    site = proj / "docs" / "index.html"
    monkeypatch.setattr(main, "ROOT", proj)
    monkeypatch.setattr(publish, "publish", lambda: site.parent.mkdir(exist_ok=True) or site.write_text(site.name + "x", encoding="utf-8"))
    monkeypatch.setenv("GITHUB_TOKEN", "github_pat_test")
    assert publish.push("Daily update 1") == "website updated on GitHub"  # a local remote: no build request
    api = "https://api.github.com/repos/callagy-recovery/litigation-tracker/pages/builds"
    asked, latest, slept = [], [{"commit": "abc123"}], []

    def urlopen(req, timeout):
        asked.append((req.get_method(), req.full_url))
        return io.BytesIO(json.dumps(latest[0]).encode() if req.get_method() == "GET" else b"")
    monkeypatch.setattr(publish.urllib.request, "urlopen", urlopen)
    monkeypatch.setattr(publish.time, "sleep", slept.append)
    remote_url = "https://github.com/callagy-recovery/litigation-tracker.git"
    # GitHub started the build of the pushed commit by itself: no second request (it would show as "errored").
    assert publish.request_build(remote_url, "t", "abc123") == "" and asked == [("GET", api + "/latest")] and slept == []
    # It never did (still showing the previous commit): wait the minute, then ask once.
    asked.clear()
    latest[0] = {"commit": "old999"}
    assert publish.request_build(remote_url, "t", "abc123") == ""
    assert asked == [("GET", api + "/latest")] * 7 + [("POST", api)] and sum(slept) == publish.BUILD_WAIT
    assert git("ls-tree", "-r", "--name-only", "main", cwd=remote).split() == [".gitignore", "docs/index.html"]
    monkeypatch.setattr(publish, "publish", lambda: None)
    assert publish.push("Daily update 2") == "website already up to date on GitHub"
    assert git("rev-list", "--count", "main", cwd=remote).strip() == "1"
    git("checkout", "-q", "-b", "local-history")
    assert "not pushed" in publish.push("Daily update 3")
    monkeypatch.delenv("GITHUB_TOKEN")
    git("checkout", "-q", "main")
    assert "no GITHUB_TOKEN" in publish.push("Daily update 4")


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js not installed")
@pytest.mark.skipif(not REPORT.exists(), reason="acceptance report not in reports\\")
def test_the_in_page_filter_matches_the_python_one(tmp_path):
    report = data.load(REPORT)
    decisions, pending = list(report["decisions"].values()), list(report["pending"].values())
    counts = {}
    for d in decisions + pending:
        for c in d["insurers"]:
            counts[c] = counts.get(c, 0) + 1
    top = sorted(counts, key=lambda c: -counts[c])[:6]
    choices = [[c] for c in top] + [top[:2], top[2:5], ["Nobody At All"]]
    cases = tmp_path / "cases.json"
    cases.write_text(json.dumps({"decisions": decisions, "pending": pending, "choices": choices,
                                 "base": data.views(decisions, pending)}), encoding="utf-8")
    script = tmp_path / "run.mjs"
    script.write_text(f"""import {{ readFileSync }} from 'node:fs';
eval(readFileSync({json.dumps(str(ROOT / 'app' / 'static' / 'views.js'))}, 'utf8'));
const c = JSON.parse(readFileSync({json.dumps(str(cases))}, 'utf8'));
console.log(JSON.stringify(c.choices.map((only) => globalThis.LTViews.views(c.decisions, c.pending, only, c.base))));
""", encoding="utf-8")
    got = json.loads(subprocess.run(["node", str(script)], capture_output=True, text=True, check=True, encoding="utf-8").stdout)
    for only, js in zip(choices, got):
        assert js == json.loads(json.dumps(data.views(decisions, pending, only))), only
