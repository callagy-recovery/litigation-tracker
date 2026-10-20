"""The acceptance checks from the brief, against NSA_Enforcement_Litigation_Dashboard_Source_2026-09-29.xlsx."""
import json
import logging
import re
import shutil

import openpyxl
import pytest

import data
import main
from conftest import REPORT

pytestmark = pytest.mark.skipif(not REPORT.exists(), reason="acceptance report not in reports\\")


@pytest.fixture(scope="module")
def report():
    return data.load(REPORT)


def coloured(view):
    return {ab: (s["decisions"], s["wins"], s["colour"], s["hatched"])
            for ab, s in view["states"].items() if s["colour"] != "grey"}


def test_row_counts(report):
    assert report["counts"]["read"] == 46
    assert report["counts"]["included"] == 43
    assert report["counts"]["excluded"] == 3
    assert [d["id"] for d in report["excludedList"]] == ["NSA-001", "NSA-026", "NSA-027"]


def test_federal_all_decisions(report):
    v = report["views"]["federal"]
    assert v["summary"]["included"] == 41
    assert coloured(v) == {
        "AZ": (1, 0, "red", True), "CA": (3, 0, "red", False), "CT": (3, 3, "green", False),
        "DE": (1, 0, "red", True), "FL": (3, 1, "orange", False), "GA": (1, 0, "red", True),
        "IL": (4, 0, "red", False), "LA": (1, 0, "red", True), "MD": (2, 2, "green", True),
        "NJ": (12, 0, "red", False), "NY": (6, 0, "red", False), "OH": (1, 0, "red", True),
        "PA": (1, 0, "red", True), "TX": (2, 0, "red", True),
    }
    assert round(v["states"]["FL"]["share"] * 100) == 33


def test_federal_appellate_only(report):
    v = report["views"]["appellate"]
    assert (v["summary"]["included"], v["summary"]["wins"], v["summary"]["losses"]) == (2, 0, 2)
    assert coloured(v) == {st: (1, 0, "red", True) for st in ("LA", "MS", "TX", "CT", "NY", "VT")}
    assert v["states"]["MS"]["via"] == "5th Circuit" and v["states"]["VT"]["via"] == "2nd Circuit"


def test_state_courts(report):
    v = report["views"]["state"]
    assert v["summary"]["included"] == 2
    assert coloured(v) == {"DE": (1, 0, "red", True), "OH": (1, 0, "red", True)}


def test_no_pending_or_dol_sheet_so_nothing_is_blue(report):
    assert report["hasPendingSheet"] is False and report["dol"] is None
    for v in report["views"].values():
        assert not any(s["colour"] == "blue" for s in v["states"].values())


def test_outcome_and_heat_map_code_agree_on_every_row(report):
    assert not [n for n in report["notes"] if n["kind"] == "Outcome and Heat Map Code disagree"]
    assert report["notes"] == []


def test_interlocutory_wins_keep_their_status(report):
    md = [report["decisions"][i] for i in report["views"]["federal"]["states"]["MD"]["ids"]]
    assert all(d["win"] and "Interlocutory" in d["status"] for d in md)
    assert [d["date"] for d in md] == sorted((d["date"] for d in md), reverse=True)  # newest first


def test_patient_initials_never_reach_the_page_or_log(tmp_path, monkeypatch):
    reports = tmp_path / "reports"  # only the acceptance report, whatever newer ones sit in reports\
    reports.mkdir()
    shutil.copy(REPORT, reports / REPORT.name)
    monkeypatch.setattr(main, "REPORTS", reports)
    monkeypatch.setattr(main, "LOGS", tmp_path)
    monkeypatch.setattr(main, "_state", {"report_sig": None, "theme": None})
    main.log.handlers.clear()
    main.setup_logging()
    try:
        status, page = main.build_page()
        main.log.info("Should be masked in the log: a/s/o X.Y.")
    finally:
        for h in main.log.handlers:
            h.close()
        main.log.handlers.clear()
    assert status == 200
    assert "a/s/o ***" in page
    assert not re.search(r"T\.\s?U\.", page)
    log_text = "".join(p.read_text(encoding="utf-8") for p in tmp_path.glob("*.log"))
    assert "Report used: NSA_Enforcement_Litigation_Dashboard_Source_2026-09-29.xlsx" in log_text
    assert "a/s/o ***" in log_text and "X.Y." not in log_text
    payload = json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>', page, re.S)[1])
    assert {k: payload["report"][k] for k in ("name", "date")} == {"name": REPORT.name, "date": "2026-09-29"}
    assert "autoRun" in payload["report"]  # the daily update's last run, for "Data as of"


def test_missing_required_sheet_or_column_is_named(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Something Else"
    p = tmp_path / "NSA_Enforcement_Litigation_Dashboard_Source_2026-10-01.xlsx"
    wb.save(p)
    with pytest.raises(data.ReportError, match='Required sheet "NSA Enforcement Decisions" is missing'):
        data.load(p)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "  nsa enforcement decisions "  # sheet and headers match after trimming, ignoring case
    ws.append([c.upper() + "  " for c in data.DECISION_COLS if c not in ("Outcome", "Court State")])
    wb.save(p)
    with pytest.raises(data.ReportError, match="missing column\\(s\\): Court State, Outcome"):
        data.load(p)


def _write_report(path, decisions, pending=None, dol=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = data.DECISIONS_SHEET
    ws.append(data.DECISION_COLS)
    for d in decisions:
        ws.append([d.get(c) for c in data.DECISION_COLS])
    for name, cols, rows in ((data.PENDING_SHEET, data.PENDING_COLS, pending), (data.DOL_SHEET, data.DOL_COLS, dol)):
        if rows is not None:
            s = wb.create_sheet(name)
            s.append(cols)
            for r in rows:
                s.append([r.get(c) for c in cols])
    wb.save(path)


def test_optional_sheets_notes_and_blue_states(tmp_path):
    p = tmp_path / "NSA_Enforcement_Litigation_Dashboard_Source_2026-10-01.xlsx"
    base = {"Court Level": "Federal District", "Circuit": "3rd", "Heat Map Include (Recommended)": "Yes"}
    _write_report(p, [
        {**base, "Decision ID": "X1", "Court State": "NJ", "Outcome": "Positive", "Heat Map Code": -1},  # disagrees
        {**base, "Decision ID": "X2", "Court State": "PA", "Outcome": "Unclear", "Heat Map Code": 0},
        {**base, "Decision ID": "X3", "Court Level": "Tribal", "Court State": "AZ", "Outcome": "Negative", "Heat Map Code": -1},
        {**base, "Decision ID": "X4", "Circuit": "Federal Cir.", "Court State": "DE", "Outcome": "Negative", "Heat Map Code": -1},
    ], pending=[{"Case Name": "Clinic a/s/o X.Y. v. Plan", "Court Level": "Federal District", "Circuit": "2d", "Court State": "NY"},
                {"Case Name": "Appeal", "Court Level": "Federal Appellate", "Circuit": "9th", "Court State": "CA"},
                {"Case Name": "State one", "Court Level": "State Trial", "Court State": "TX"}],
        dol=[{"Title": "Old rule", "Date Issued": "2024-01-01"}, {"Title": "New rule", "Date Issued": "2026-08-01", "Status": "Final"}])
    r = data.load(p)
    kinds = {n["kind"] for n in r["notes"]}
    assert {"Outcome and Heat Map Code disagree", "Unknown Outcome", "Unknown Court Level", "Unknown Circuit"} <= kinds
    assert r["views"]["federal"]["states"]["NJ"]["colour"] == "green"  # Outcome wins over Heat Map Code
    assert r["views"]["federal"]["states"]["PA"]["decisions"] == 0      # unknown outcome is not counted
    assert r["views"]["federal"]["states"]["NY"]["colour"] == "blue"
    assert r["views"]["state"]["states"]["TX"]["colour"] == "blue"
    assert r["views"]["state"]["states"]["NY"]["colour"] == "grey"
    appellate = r["views"]["appellate"]["states"]
    assert all(appellate[s]["colour"] == "blue" for s in data.CIRCUIT_STATES["9"])
    assert r["pending"]["P2"]["case"] == "Clinic a/s/o *** v. Plan"
    assert [a["title"] for a in r["dol"]] == ["New rule", "Old rule"]
    assert r["views"]["federal"]["summary"]["pending"] == 2



def test_insurer_filter_and_case_lists(tmp_path):
    p = tmp_path / "NSA_Enforcement_Litigation_Dashboard_Source_2026-10-01.xlsx"
    base = {"Court Level": "Federal District", "Circuit": "3rd", "Heat Map Include (Recommended)": "Yes", "Court State": "NJ"}
    _write_report(p, [
        {**base, "Decision ID": "A", "Outcome": "Positive", "Heat Map Code": 1, "Decision Date": "2026-01-02", "Payor / Defendant": "AETNA LIFE INSURANCE COMPANY"},
        {**base, "Decision ID": "B", "Outcome": "Negative", "Heat Map Code": -1, "Decision Date": "2026-03-04", "Payor / Defendant": "Aetna Health / Cigna entities"},
        {**base, "Decision ID": "C", "Outcome": "Negative", "Heat Map Code": -1, "Decision Date": "2025-05-06", "Payor / Defendant": "Horizon BCBS of New Jersey"},
    ], pending=[{"Case Name": "P v. Q", "Court Level": "Federal District", "Circuit": "3d", "Court State": "NJ", "Date Filed": "2026-06-01",
                 "Payor / Defendant": "UMR, Inc."}])
    r = data.load(p)
    assert r["decisions"]["B"]["insurers"] == ["Aetna (CVS Health)", "Cigna"]
    fed = r["views"]["federal"]
    assert fed["caseIds"] == ["B", "A", "C"] and fed["pendingIds"] == ["P2"]  # newest first
    only = data.views(list(r["decisions"].values()), list(r["pending"].values()), ["Aetna (CVS Health)"])["federal"]
    assert only["caseIds"] == ["B", "A"] and only["pendingIds"] == [] and only["states"]["NJ"]["decisions"] == 2
    both = data.views(list(r["decisions"].values()), list(r["pending"].values()), ["Cigna", "UnitedHealth Group"])["federal"]
    assert both["caseIds"] == ["B"] and both["pendingIds"] == ["P2"]
