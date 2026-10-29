"""The daily update: the two-level reader's rules (with stand-in readers, no Claude calls) and how the dashboard merges
what it found with the report."""
import datetime as dt
import json
import os
import urllib.parse

import pytest

import data
import reader
from test_acceptance import _write_report

ORDER = ("The provider petitioned to confirm the IDR award against the plan. " * 10 +
         "For these reasons, the petition to confirm the award is GRANTED and judgment is entered for the provider.")
QUOTE = "For these reasons, the petition to confirm the award is GRANTED and judgment is entered for the provider."


def answer(result="Win", quote=QUOTE, relevant=True):
    return {"relevant": relevant, "result": result, "final": True, "quote": quote, "reason": "r", "holding": "h"}


def scripted(l1, l2, calls=None):
    """A stand-in ask(): Level 1 and Level 2 give their scripted answers in turn and record each prompt."""
    seq = {reader.L1: iter(l1), reader.L2: iter(l2)}

    def ask(model, system, prompt):
        if calls is not None:
            calls.append((model, prompt))
        return next(seq[model])
    return ask


def test_both_readers_agree_first_time_and_read_blind():
    calls = []
    r = reader.read(ORDER, "P v. Q", "2026-10-01", "rules", ask=scripted([answer()], [answer()], calls))
    assert r["agreed"] and r["result"] == "Win" and r["rounds"] == 1
    assert len(calls) == 2 and "Another reader" not in calls[1][1]  # Level 2 never sees Level 1's answer in round 1


def test_disagreement_goes_back_and_forth_then_stays_pending():
    calls = []
    r = reader.read(ORDER, "P v. Q", "2026-10-01", "rules", ask=scripted([answer("Win")] * 3, [answer("Loss")] * 3, calls))
    assert not r["agreed"] and r["rounds"] == reader.ROUNDS and "did not agree" in r["note"]
    assert len(calls) == 2 * reader.ROUNDS and "Another reader concluded result=Loss" in calls[2][1]


def test_a_made_up_quote_is_sent_back():
    r = reader.read(ORDER, "P v. Q", "2026-10-01", "rules",
                    ask=scripted([answer(quote="The court enforces every award ever made."), answer()], [answer(), answer()]))
    assert r["agreed"] and r["rounds"] == 2


def test_an_answer_against_the_docket_wording_is_not_accepted():
    r = reader.read(ORDER, "P v. Q", "2026-10-01", "rules", "Ruling: petition to confirm denied", "ORDER denying petition",
                    ask=scripted([answer("Win")] * 3, [answer("Win")] * 3))
    assert not r["agreed"] and "contradicts the docket entry" in r["note"]


def test_no_readable_text_is_never_guessed():
    r = reader.read("Doc. 13", "P v. Q", "2026-10-01", "rules", ask=scripted([], []))
    assert not r["agreed"] and r["rounds"] == 0


def _merge_setup(tmp_path, found):
    p = tmp_path / "NSA_Enforcement_Litigation_Dashboard_Source_2026-10-01.xlsx"
    base = {"Court Level": "Federal District", "Circuit": "3rd", "Court State": "NJ", "Heat Map Include (Recommended)": "Yes"}
    _write_report(p, [{**base, "Decision ID": "A", "Outcome": "Negative", "Heat Map Code": -1, "Decision Date": "2026-03-04",
                       "Docket No.": "2:25-cv-00100", "Case Name": "Clinic v. Plan"}],
                  pending=[{"Case Name": "Doc a/s/o X.Y. v. Horizon", "Court Level": "Federal District", "Circuit": "3d",
                            "Court State": "NJ", "Docket No.": "2:25-cv-00200", "Date Filed": "2025-06-01"},
                           {"Case Name": "Settled v. Aetna", "Court Level": "Federal District", "Circuit": "3d",
                            "Court State": "NJ", "Docket No.": "2:25-cv-00300", "Date Filed": "2025-06-02"}], dol=[])
    auto = tmp_path / "auto"
    auto.mkdir()
    (auto / data.AUTO_FILE).write_text(json.dumps(found), encoding="utf-8")
    return data.load(p, auto)


def test_merge_report_wins_and_rulings_replace_pending(tmp_path):
    today = dt.date.today().isoformat()
    ruling = {"level": "Federal District", "circuit": "3", "state": "NJ", "result": "Win", "final": True, "holding": "Confirmed."}
    r = _merge_setup(tmp_path, {"lastRun": today, "decisions": [
        {**ruling, "case": "Clinic v. Plan", "docket": "25-cv-100", "date": "2026-03-05", "result": "Loss"},  # report has it
        {**ruling, "case": "Doc a/s/o X.Y. v. Horizon", "docket": "2:25-cv-00200-KM", "date": "2026-10-02",
         "payer": "Horizon BCBS of New Jersey", "url": "https://example.org/o.pdf"}],
        "ended": [{"docket": "25-cv-300", "circuit": "3", "state": "NJ"}],
        "pending": [{"case": "New a/s/o J.K. v. Cigna", "level": "Federal District", "circuit": "3", "state": "NJ",
                     "docket": "26-cv-5", "filed": "2026-10-01", "payer": "Cigna"}],
        "dol": [{"title": "New IDR rule", "issued": "2026-10-01"}]})
    auto = [d for d in r["decisions"].values() if d.get("auto")]
    assert [d["case"] for d in auto] == ["Doc a/s/o *** v. Horizon"]  # the duplicate of report row A is dropped; masked
    assert auto[0]["win"] and auto[0]["included"] and auto[0]["insurers"] == ["Horizon Blue Cross Blue Shield of New Jersey"]
    assert [p["case"] for p in r["pending"].values()] == ["New a/s/o *** v. Cigna"]  # decided and ended rows are gone
    assert r["views"]["federal"]["states"]["NJ"]["decisions"] == 2 and r["views"]["federal"]["states"]["NJ"]["wins"] == 1
    assert [a["title"] for a in r["dol"]] == ["New IDR rule"] and r["dol"][0]["auto"]
    kinds = {n["kind"] for n in r["notes"]}
    assert "Automatic update" in kinds
    assert "Automatic updates have stopped" not in kinds
    assert r["autoRun"] == today  # the page's "Data as of" uses it
    assert "X.Y" not in json.dumps(r) and "J.K" not in json.dumps(r)


def test_a_companion_docket_decided_the_same_day_is_its_own_ruling(tmp_path):
    # The editor codes each docket as its own ruling (NSA-042 and NSA-048 are one judge's same-day orders).
    r = _merge_setup(tmp_path, {"lastRun": dt.date.today().isoformat(), "decisions": [
        {"case": "Clinic, LLC v. Plan Inc.", "docket": "3:25-cv-01367", "date": "2026-03-04", "level": "Federal District",
         "circuit": "3", "state": "NJ", "result": "Loss"}]})
    assert [d["docket"] for d in r["decisions"].values() if d.get("auto")] == ["3:25-cv-01367"]


FILING = "The plaintiff provider petitions this Court to confirm the IDR award issued against the defendant plan. " * 8
GOOD = {"enforcement": True, "quote": "petitions this Court to confirm the IDR award issued against the defendant plan", "reason": "r"}


def test_screen_needs_both_readers_and_real_quotes():
    calls = []
    assert reader.screen(FILING, "P v. Q", ask=scripted([GOOD], [GOOD], calls)) is True
    assert [m for m, _ in calls] == [reader.L1, reader.L2] and "Another reader" not in calls[1][1]  # Opus reads blind
    bad = {**GOOD, "quote": "a sentence the filing never says at all, invented"}
    assert reader.screen(FILING, "P v. Q", ask=lambda m, s, p: bad) is None
    assert reader.screen("Doc. 1", "P v. Q", ask=lambda m, s, p: GOOD) is None


def test_screen_disagreement_goes_back_and_forth():
    no = {**GOOD, "enforcement": False}
    assert reader.screen(FILING, "P v. Q", ask=scripted([GOOD] * 3, [no] * 3)) is None  # never agree: check by hand
    calls = []
    assert reader.screen(FILING, "P v. Q", ask=scripted([GOOD, GOOD], [no, GOOD], calls)) is True  # Opus changes its mind
    assert "Another reader concluded enforcement=True" in calls[3][1]


def test_party_filter_keeps_provider_suits_only():
    import daily_update
    assert daily_update.provider_v_payer("GuardION Medical LLC v. MedCost Benefit Services LLC")
    assert daily_update.provider_v_payer("Some Clinic P.C. v. Aetna Life Insurance Co.")
    assert not daily_update.provider_v_payer("Blue Cross Blue Shield Healthcare Plan of Georgia, Inc. v. HaloMD, Inc.")
    assert not daily_update.provider_v_payer("Texas Medical Association v. HHS")


def test_dol_items_must_be_closely_related():
    import daily_update
    assert daily_update.relevant_dol({"title": "Federal Independent Dispute Resolution Operations; Correction"})
    assert daily_update.relevant_dol({"title": "Agency Information Collection", "abstract": "IDR process forms"})
    assert not daily_update.relevant_dol({"title": "Excepted Fertility Benefits", "abstract": "... the No Surprises Act ..."})


CASE = {"case": "Alpha Med v. Aetna", "court": "District Court, D. New Jersey", "level": "Federal District", "circuit": "3",
        "state": "NJ", "docket": "2:26-cv-00001", "url": "https://www.courtlistener.com/docket/1/x/", "payer": "Aetna"}


def test_what_the_readers_cannot_confirm_goes_on_the_to_check_list(monkeypatch):
    import daily_update as du
    monkeypatch.setattr(du, "order_text", lambda api, token, c: "order text")
    answers = {"Agreed v. Plan": {"agreed": True, "relevant": True, "result": "Win", "final": True, "holding": "Confirmed.", "rounds": 1},
               "Disagreed v. Plan": {"agreed": False, "rounds": 3, "note": "Could not agree after 3 rounds."},
               "Decided v. Plan": {"agreed": True, "relevant": True, "result": "Loss", "final": True, "holding": "Reversed.", "rounds": 1}}
    monkeypatch.setattr(reader, "read", lambda text, case, *a: answers[case])
    queue = [{**CASE, "case": c, "key": c, "date": "2026-10-01", "fetch": [("pdf", "x")], "decided": c.startswith("Decided")}
             for c in answers]
    found, state, asks = {"decisions": []}, {"seen": {}}, []
    du.read_queue(None, "", queue, state, found, "rules", 10, asks)
    assert [d["case"] for d in found["decisions"]] == ["Agreed v. Plan"]  # confirmed: straight on the dashboard
    assert {a["case"]: a["kind"] for a in asks} == {"Disagreed v. Plan": "Ruling", "Decided v. Plan": "Later ruling in a decided case"}
    assert "Loss" in next(a for a in asks if a["case"] == "Decided v. Plan")["readers"]
    assert queue == []


def test_answers_on_the_to_check_list_reach_the_dashboard(tmp_path):
    import openpyxl
    import daily_update as du
    import to_check as tc
    store = {}
    items = [tc.item("docket:11", "New case", "The two readers could not tell.", {**CASE, "docketId": "11", "date": "2026-10-02"}),
             tc.item("rd:22", "Ruling", "Could not agree after 3 rounds.", {**CASE, "case": "Beta ER v. Cigna", "docket": "2:26-cv-00002",
                                                                            "date": "2026-10-05"}),
             tc.item("docket:33", "New case", "No free copy of the first filing.", {**CASE, "case": "Gamma v. Anthem", "docketId": "33",
                                                                                    "docket": "2:26-cv-00003"})]
    f, n = tc.add(tmp_path, store, items, dt.date(2026, 10, 20))
    assert n == 3 and f.name == "To check.xlsx"
    assert tc.add(tmp_path, store, items, dt.date(2026, 10, 21)) == (None, 0)  # never asked twice

    wb = openpyxl.load_workbook(f)
    ws = wb["To check"]
    ws["J2"], ws["J3"], ws["K3"] = "Pending case", "win", "Petition to confirm granted."
    wb.save(f)
    assert sorted(tc.collect(tmp_path, store)) == ["docket:11", "rd:22"]
    assert tc.tidy(tmp_path, store) == 2  # answered rows leave the To check sheet for the Settled sheet
    wb = openpyxl.load_workbook(f)
    assert [r[0].value for r in wb["To check"].iter_rows(min_row=2)] == ["Gamma v. Anthem"]
    assert [(r[0].value, r[9].value, r[14].value) for r in wb["Settled"].iter_rows(min_row=2)] == [
        ("Alpha Med v. Aetna", "Pending case", "a person"), ("Beta ER v. Cigna", "Win", "a person")]
    assert tc.open_count(store) == 1
    # A later run adds its rows at the bottom of the same file, with the day they were found.
    later = tc.item("docket:44", "New case", "No free copy of the first filing.", {**CASE, "case": "Delta v. Cigna", "docketId": "44"})
    assert tc.add(tmp_path, store, [later], dt.date(2026, 10, 23)) == (f, 1)
    wb = openpyxl.load_workbook(f)
    assert [(r[0].value, r[2].value, r[12].value) for r in wb["To check"].iter_rows(min_row=2)] == [
        ("Gamma v. Anthem", "2:26-cv-00003", "2026-10-20"), ("Delta v. Cigna", CASE["docket"], "2026-10-23")]
    ws = wb["To check"]
    ws["J2"], ws["J3"] = "Leave out", "Leave out"
    wb.save(f)
    tc.collect(tmp_path, store)
    tc.tidy(tmp_path, store)
    assert f.exists() and not (tmp_path / tc.FOLDER / tc.DONE).exists()  # every row settled: the one file stays
    assert [r[0].value for r in openpyxl.load_workbook(f)["To check"].iter_rows(min_row=2)] == []

    state = {"dockets": {"11": {"enforcement": None}, "33": {"enforcement": True}}}
    found = {"decisions": [], "pending": [{**CASE, "case": "Gamma v. Anthem", "docketId": "33"}]}
    assert du.apply_answers(store, state, found) == (2, 2)  # Gamma and Delta left out
    assert du.apply_answers(store, state, found) == (2, 2)  # the same answers again: not added twice
    assert [p["case"] for p in found["pending"]] == ["Alpha Med v. Aetna"]  # Gamma left out, Alpha confirmed by hand
    assert found["decisions"][0]["result"] == "Win" and found["decisions"][0]["holding"] == "Petition to confirm granted."
    assert state["dockets"]["11"]["enforcement"] is True and state["dockets"]["33"]["enforcement"] is False

    (tmp_path / "auto").mkdir(exist_ok=True)
    (tmp_path / "auto" / data.AUTO_FILE).write_text(json.dumps(found), encoding="utf-8")
    decisions, pending = [], []
    data.merge_auto(tmp_path / "auto", decisions, pending, None, [])
    assert decisions[0]["byHand"] and decisions[0]["status"] == "Checked by a person" and decisions[0]["caution"] == data.HAND_CAUTION
    assert pending[0]["byHand"] and pending[0]["auto"]


def test_a_scanned_pdf_is_copied_out_page_by_page(tmp_path):
    from PIL import Image, ImageDraw
    import daily_update
    pages = []
    for n in range(2):  # a PDF that is only pictures of pages, like a court scan
        im = Image.new("L", (850, 1100), 255)
        ImageDraw.Draw(im).text((50, 50), f"Page {n + 1}: the petition to confirm is GRANTED.", fill=0)
        pages.append(im.convert("RGB"))
    scan = tmp_path / "scan.pdf"
    pages[0].save(scan, save_all=True, append_images=pages[1:])
    seen = []
    text = daily_update.file_text(scan, transcribe=lambda png: seen.append(png[:8]) or "copied page text")
    assert text == "copied page text\n\ncopied page text" and all(s == b"\x89PNG\r\n\x1a\n" for s in seen)
    assert daily_update.file_text(tmp_path / "missing.pdf") == ""


def summary_ask(writes, checks, calls=None):
    """A stand-in for the summary writer (Sonnet) and checker (Opus)."""
    w, c = iter(writes), iter(checks)

    def ask(model, system, prompt, schema):
        if calls is not None:
            calls.append((model, prompt))
        return next(w) if model == reader.L1 else next(c)
    return ask


SUMMARY = {"summary": "The court confirmed the award.", "result": "Win", "quote": QUOTE, "conflict": ""}


def test_summary_needs_the_checker_and_the_coded_result():
    calls = []
    r = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Win", ask=summary_ask([SUMMARY], [{"ok": True, "problems": ""}], calls))
    assert r["agreed"] and r["summary"] == "The court confirmed the award." and [m for m, _ in calls] == [reader.L1, reader.L2]
    fixed = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Win", ask=summary_ask(
        [SUMMARY, {**SUMMARY, "summary": "Fixed."}], [{"ok": False, "problems": "It also denied fees."}, {"ok": True, "problems": ""}], calls))
    assert fixed["agreed"] and fixed["summary"] == "Fixed." and fixed["rounds"] == 2 and "denied fees" in calls[-2][1]
    wrong = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Loss", ask=summary_ask([SUMMARY] * 3, []))
    assert not wrong["agreed"]  # the summary must match the editor's coded result
    conflict = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Loss", ask=summary_ask([{**SUMMARY, "conflict": "It was granted."}], []))
    assert not conflict["agreed"] and "may not support" in conflict["note"]
    made_up = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Win", ask=summary_ask([{**SUMMARY, "quote": "Never said in this order at all."}] * 3, []))
    assert not made_up["agreed"]


def test_summary_with_a_patients_initials_is_rewritten():
    assert reader.initials_in("Northeast moved only on L.T.'s claim.") == "L.T."
    assert reader.initials_in("The U.S. District Court (D.C.) ruled for Shareef Jandali, LLC and Clinic P.C. in N.J.") is None
    r = reader.summarise(ORDER, "P v. Q", "2026-10-01", "Win", ask=summary_ask(
        [{**SUMMARY, "summary": "Confirmed J. D.'s award."}, SUMMARY], [{"ok": True, "problems": ""}]))
    assert r["agreed"] and r["rounds"] == 2 and "J." not in r["summary"]


def test_summaries_show_only_for_the_outcome_they_were_written_for(tmp_path):
    decisions = [{"docket": "3:25-cv-01367", "date": "2026-07-21", "win": True, "counted": True},
                 {"docket": "2:25-cv-00100", "date": "2026-03-04", "win": False, "counted": True},
                 {"docket": "1:26-cv-00009", "date": "2026-05-05", "win": True, "counted": True}]
    (tmp_path / data.SUMMARY_FILE).write_text(json.dumps({
        "25cv1367|2026-07-21": {"outcome": "Win", "summary": "Confirmed a/s/o X.Y.'s award."},
        "25cv100|2026-03-04": {"outcome": "Win", "summary": "Written when it was coded a win."},  # the report now says Loss
        "26cv9|2026-05-05": {"outcome": "Win", "failed": "No agreed summary."}}), encoding="utf-8")
    data.merge_summaries(tmp_path, decisions)
    assert decisions[0]["summary"] == "Confirmed a/s/o ***'s award." and "summary" not in decisions[1] and "summary" not in decisions[2]


def test_a_missing_ruling_copy_skips_only_that_ruling(tmp_path, monkeypatch):
    import daily_update as du
    monkeypatch.setattr(du, "AUTO", tmp_path)
    monkeypatch.setattr(du, "SUMMARY_SOURCES", tmp_path / "summary_sources.json")
    pdf = "https://storage.courtlistener.com/recap/gov.uscourts.njd.1/gov.uscourts.njd.1.{}.0.pdf"
    decisions = {f"N{n}": {"docket": f"2:25-cv-0010{n}", "date": f"2026-09-0{n}", "case": f"Clinic {n} v. Plan", "win": False,
                           "url": pdf.format(n)} for n in (1, 2, 3)}
    report = {"decisions": decisions, "views": {v: {"caseIds": ["N3", "N2", "N1"]} for v in ("state", "federal", "appellate")}}
    errors = {pdf.format(3): "HTTP 404 from www.courtlistener.com"}  # the newest one's page is gone

    def text(url):
        if url in errors:
            raise RuntimeError(errors[url])
        return "Clinic v. Plan: order text"
    monkeypatch.setattr(du, "ruling_text", text)
    monkeypatch.setattr(reader, "summarise", lambda *a: {"agreed": True, "summary": "Dismissed."})
    assert du.write_summaries(report) == {"written": 2, "could not agree": 0, "no readable copy": 1}
    errors[pdf.format(1)] = errors[pdf.format(2)] = "HTTP 429 from www.courtlistener.com"
    (tmp_path / data.SUMMARY_FILE).unlink()
    assert du.write_summaries(report) == {"written": 0, "could not agree": 0, "no readable copy": 1}  # 429: stop for today


def test_a_courtlistener_page_number_is_the_case_not_the_document(monkeypatch):
    # The REACH Air (11th Cir.) page is /opinion/10738691/; opinion document 10738691 is an unrelated Maine order.
    import daily_update as du
    import find_updates as fu
    asked = []

    def get(self, url, headers=None, body=None, cl=False):
        asked.append(url.split("?")[0])
        if "/clusters/" in url:
            return {"sub_opinions": ["https://www.courtlistener.com/api/rest/v4/opinions/11205276/"]}
        return {"plain_text": "REACH Air Medical Services v. Kaiser Foundation Health Plan. " + "affirmed " * 80}
    monkeypatch.setattr(fu.Api, "get", get)
    text = du.ruling_text("https://www.courtlistener.com/opinion/10738691/reach-air-medical-services-llc-v-kaiser/")
    assert asked == ["https://www.courtlistener.com/api/rest/v4/clusters/10738691/",
                     "https://www.courtlistener.com/api/rest/v4/opinions/11205276/"] and text.startswith("REACH Air")
    case = "REACH Air Medical Services LLC v. Kaiser Foundation Health Plan Inc."
    assert du.same_case(text, case)
    assert not du.same_case("UNITED STATES DISTRICT COURT DISTRICT OF MAINE STEPHEN J. RATCLIFFE v. BRP U.S., INC.", case)
    assert du.same_case("Petition granted.", "AA Med., P.C. v. Aetna Life Insurance Co.") is False  # Aetna not named
    assert du.same_case("anything", "Med. v. Inc.")  # no name to check: not held back


def test_only_links_the_update_can_read_are_used_for_summaries():
    import daily_update
    assert daily_update.readable("https://storage.courtlistener.com/recap/gov.uscourts.ctd.1/gov.uscourts.ctd.1.25.0.pdf")
    assert not daily_update.readable("https://law.justia.com/cases/federal/district-courts/x/1/")
    assert not daily_update.readable("https://www.courtlistener.com/docket/70253986/worldwide/")


def test_newsletter_drafts_at_most_daily_and_only_when_there_is_something_new(tmp_path):
    import newsletter
    ruling = lambda n, win=False: {"id": f"N{n}", "docket": f"2:25-cv-{n:05d}", "date": f"2026-10-{n:02d}", "case": f"Clinic {n} v. Plan",  # noqa: E731
                                   "court": "D.N.J.", "state": "NJ", "win": win, "included": True, "counted": True,
                                   "status": "Final district ruling", "holding": "Dismissed.", "url": "https://example.org/o.pdf"}
    report = {"decisions": {"N1": ruling(1)}, "pending": {}}
    drafts, said = [], []
    make = lambda s, b, bcc, logo=None: drafts.append((s, b, bcc))  # noqa: E731
    day = dt.date(2026, 10, 16)
    auto = tmp_path / "auto"
    auto.mkdir()
    assert newsletter.run(report, "r.xlsx", auto, said.append, day, make).startswith("started")  # first run: no draft
    report["decisions"]["N2"] = {**ruling(2, win=True), "summary": "The court confirmed the award."}
    assert newsletter.run(report, "r.xlsx", auto, said.append, day, make).startswith("not due")  # at most one a day
    import signups
    signups.subscribe(tmp_path, "pat@callagylaw.com")
    signups.subscribe(tmp_path, "sam@callagyrecovery.com")
    signups.subscribe(tmp_path, "gone@callagylaw.com")
    signups.unsubscribe(tmp_path, "gone@callagylaw.com")
    report["pending"] = {f"P{n}": {"case": f"Doc {n} v. Aetna", "court": "D.N.J.", "state": "NJ" if n < 15 else "NY",
                                   "docket": f"3:26-cv-{n}", "filed": "2026-10-10"} for n in range(17)}
    assert newsletter.run(report, "r.xlsx", auto, said.append, day + dt.timedelta(days=1), make).startswith("draft saved")
    subject, body, bcc = drafts[0]
    assert subject == "Litigation Tracker · Oct 17, 2026: 1 new ruling, 17 new cases" and bcc == "pat@callagylaw.com; sam@callagyrecovery.com"  # everyone signed up, in BCC
    assert "Update · Saturday, Oct 17, 2026" in body and "updated for Oct 17, 2026." in body and "Oct 16" not in body
    assert "Clinic 2 v. Plan" in body and ">WIN<" in body and "Clinic 1 v. Plan" not in body
    assert ">NJ<" in body and ">15</td>" in body and ">2</td>" in body and "Doc 3" not in body  # cases are counted by state
    assert "The court confirmed the award." not in body and "dashboard" in body  # brief: the details are on the dashboard
    assert f"href='{newsletter.SITE_URL}'" in body and ">Click Here</a>" in body and "Or open <b>Litigation Tracker</b> from Callagy Hub" in body
    assert "width='600'" in body  # email-safe 600 px layout
    assert newsletter.short_name("Shareef Jandali Plastic Surgery, LLC v. Cigna Health and Life Insurance Co.", "Cigna") == \
        "Shareef Jandali Plastic Surgery v. Cigna"
    assert newsletter.short_name("Jeffrey Farkas, M.D., LLC v. Aetna Life Insurance Company", "Aetna (CVS Health)") == "Jeffrey Farkas v. Aetna"
    assert newsletter.run(report, "r.xlsx", auto, said.append, day + dt.timedelta(days=2), make) == "due, but nothing new: no draft"
    assert len(drafts) == 1


def test_a_later_run_the_same_day_adds_to_that_days_draft_while_it_is_unsent(tmp_path):
    import newsletter
    pend = lambda n, st="NJ": {"case": f"Doc {n} v. Aetna", "court": "D.N.J.", "state": st, "docket": f"3:26-cv-{n}", "filed": "2026-10-10"}  # noqa: E731
    report = {"decisions": {}, "pending": {"P1": pend(1)}}
    auto, said, made, updates = tmp_path / "auto", [], [], []
    auto.mkdir()
    make = lambda s, b, bcc, logo=None: made.append(s) or {"entryId": "E1", "bodyHash": "H1"}  # noqa: E731
    outcome = ["updated"]

    def update(entry_id, body_hash, s, b, bcc=""):
        updates.append((entry_id, body_hash, s, b))
        return outcome[0], ("H2" if outcome[0] == "updated" else "")
    day = dt.date(2026, 10, 28)
    run = lambda d: newsletter.run(report, "r.xlsx", auto, said.append, d, make, update)  # noqa: E731
    run(day - dt.timedelta(days=1))  # first run: baseline
    report["pending"]["P2"] = pend(2)
    assert run(day).startswith("draft saved") and made == ["Litigation Tracker · Oct 28, 2026: 1 new case"]
    assert run(day) == "today's draft is up to date" and updates == []  # nothing new: the draft is not touched
    report["pending"]["P3"] = pend(3, "CA")  # found by a second run today (Udani, Oct 28)
    assert run(day) == "today's draft updated (+0 rulings, +1 cases)" and len(made) == 1  # same draft, no second one
    entry_id, body_hash, subject, body = updates[-1]
    assert (entry_id, body_hash) == ("E1", "H1") and subject == "Litigation Tracker · Oct 28, 2026: 2 new cases"
    assert ">CA<" in body and ">NJ<" in body and "updated for Oct 28, 2026." in body  # the morning's case kept
    report["pending"]["P4"] = pend(4)
    outcome[0] = "sent"  # sent in the meantime: left alone, and the new case waits for tomorrow
    assert run(day) == "today's draft was already sent; the 1 new items go in tomorrow's draft" and updates[-1][1] == "H2"
    outcome[0] = "edited"
    assert run(day).startswith("today's draft was edited by hand, so it was left as it is")
    assert run(day + dt.timedelta(days=1)).startswith("draft saved") and made[-1].endswith("1 new case")  # P4 tomorrow


def test_signup_list_keeps_one_row_per_person(tmp_path):
    import signups
    day = dt.date(2026, 10, 16)
    assert signups.subscribe(tmp_path, "Jane.Doe@CallagyRecovery.com", "Jane Doe", "jane.doe", day) == "added"
    assert signups.subscribe(tmp_path, " jane.doe@callagyrecovery.com ", today=day) == "already"  # capitals and spaces ignored
    assert signups.subscribe(tmp_path, "sam@callagylaw.com", today=day) == "added"
    assert signups.unsubscribe(tmp_path, "SAM@callagylaw.com", today=day) == "removed"
    assert signups.recipients(tmp_path) == ["jane.doe@callagyrecovery.com"]
    assert signups.subscribe(tmp_path, "sam@callagylaw.com", today=day) == "resubscribed"
    assert len(signups._rows(signups.path(tmp_path))) == 2  # still one row each
    for bad in ("someone@gmail.com", "not an email", ""):
        with pytest.raises(signups.SignupError):
            signups.subscribe(tmp_path, bad)
    assert signups.guess("jane.doe") == ("jane.doe@callagyrecovery.com", "Jane Doe")


def test_a_refused_document_is_skipped_not_the_whole_run():
    import daily_update
    blocked = ('Claude stopped (r model.\n\nLearn more: https://support.claude.com/en/articles/8106465\n\nDetails: `[bio]`)')
    assert isinstance(reader._failure(blocked), reader.Declined)
    assert not isinstance(reader._failure("Claude stopped (You've hit your usage limit; it resets at 3pm)"), reader.Declined)
    assert daily_update.main_document("https://storage.courtlistener.com/recap/gov.uscourts.moed.231846/gov.uscourts.moed.231846.1.2.pdf") \
        == "https://storage.courtlistener.com/recap/gov.uscourts.moed.231846/gov.uscourts.moed.231846.1.0.pdf"
    assert daily_update.main_document("https://x/gov.uscourts.njd.1.25.0.pdf") == "https://x/gov.uscourts.njd.1.25.0.pdf"


def test_merge_warns_when_the_update_has_stopped(tmp_path):
    r = _merge_setup(tmp_path, {"lastRun": (dt.date.today() - dt.timedelta(days=9)).isoformat()})
    assert "Automatic updates have stopped" in {n["kind"] for n in r["notes"]}


def test_no_auto_folder_means_the_report_alone(tmp_path):
    p = tmp_path / "NSA_Enforcement_Litigation_Dashboard_Source_2026-10-01.xlsx"
    _write_report(p, [{"Decision ID": "A", "Court Level": "Federal District", "Circuit": "3rd", "Court State": "NJ",
                       "Outcome": "Negative", "Heat Map Code": -1, "Heat Map Include (Recommended)": "Yes"}])
    r = data.load(p, tmp_path / "auto")
    assert list(r["decisions"]) == ["A"] and not any(n["kind"].startswith("Automatic") for n in r["notes"])


def test_state_cases_saved_by_hand_are_read_and_shown(tmp_path, monkeypatch):
    import daily_update  # noqa: F401  (puts app\tools on the path)
    import openpyxl
    import state_cases as sc
    import to_check as tc
    texts = {"ruling": "ORDER. The petition to confirm is granted. ", "complaint": "PETITION. We ask the court to confirm the IDR award. ",
             "brief": "BRIEF in opposition. ", "unclear": "Something hard to read. "}
    details = {"ruling": {"agreed": True, "case": "Alpha Med v. Aetna", "court": "Supreme Court, Kings County", "level": "State Trial",
                          "state": "NY", "docket": "512345/2026", "date": "2026-10-12", "kind": "ruling"},
               "complaint": {"agreed": True, "case": "Beta ER v. Cigna", "court": "Superior Court, Morris County", "level": "State Trial",
                             "state": "NJ", "docket": "MRS-L-1-26", "date": "2026-10-01", "kind": "complaint"},
               "brief": {"agreed": True, "case": "Gamma v. Anthem", "court": "Circuit Court, Miami-Dade", "level": "State Trial",
                         "state": "FL", "docket": "2026-1", "date": "2026-10-02", "kind": "other"},
               "unclear": {"agreed": False, "note": "The two readers did not agree on what this document is.", "case": "Delta v. Humana",
                           "court": "", "level": "State Trial", "state": "TX", "docket": "", "date": ""}}
    kind_of = lambda text: next(k for k, t in texts.items() if text.startswith(t))  # noqa: E731
    monkeypatch.setattr(reader, "identify", lambda text: details[kind_of(text)])
    monkeypatch.setattr(reader, "screen", lambda text, case: True)
    monkeypatch.setattr(reader, "read", lambda text, case, date, system: {
        "agreed": True, "relevant": True, "result": "Win", "final": True, "holding": "Award confirmed.", "rounds": 1})
    sc.make_folders(tmp_path)
    new = tmp_path / sc.FOLDER / sc.NEW
    for k in texts:
        (new / f"{k}.pdf").write_bytes(b"%PDF-")
    (new / "notes.txt").write_text("not a court document", encoding="utf-8")  # ignored: not a PDF or web page
    files = {f"{k}.pdf": t * 20 for k, t in texts.items()}
    found, state, summaries = {}, {"stateCases": {}}, []
    counts, asks = sc.run(tmp_path, found, state, "rules", lambda f: files[f.name], lambda m: None, lambda *a: summaries.append(a))
    assert counts == {"read": 4, "rulings": 1, "pending": 1, "to check": 2, "waiting": 0}
    assert found["decisions"][0]["case"] == "Alpha Med v. Aetna" and found["decisions"][0]["state"] == "NY"
    assert found["pending"][0]["case"] == "Beta ER v. Cigna" and found["pending"][0]["filed"] == "2026-10-01"
    assert len(summaries) == 1 and summaries[0][3] == "Win"
    assert sorted(a["case"] for a in asks) == ["Delta v. Humana", "Gamma v. Anthem"]
    assert [f.name for f in new.iterdir()] == ["notes.txt"]  # every read file moved to Done
    assert len(list((tmp_path / sc.FOLDER / sc.DONE).iterdir())) == 4 and "stateCases" not in state

    # It shows on the dashboard's State courts tab.
    (tmp_path / "auto").mkdir()
    (tmp_path / "auto" / data.AUTO_FILE).write_text(json.dumps(found), encoding="utf-8")
    decisions, pending = [], []
    data.merge_auto(tmp_path / "auto", decisions, pending, None, [])
    assert decisions[0]["level"] == "state" and decisions[0]["auto"] and pending[0]["level"] == "state"

    # On the To check list the person corrects what the readers got wrong, then answers.
    store = {}
    f, _ = tc.add(tmp_path, store, asks)
    wb = openpyxl.load_workbook(f)
    ws = wb["To check"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Delta v. Humana")
    ws.cell(row, 2).value, ws.cell(row, 5).value, ws.cell(row, 10).value = "District Court, Harris County", "2026-10-05", "Win"
    wb.save(f)
    tc.collect(tmp_path, store)
    win = tc.dashboard_rows(store)[0]["decisions"][0]
    assert (win["case"], win["court"], win["state"], win["date"]) == ("Delta v. Humana", "District Court, Harris County", "TX", "2026-10-05")


def test_both_readers_must_agree_on_what_a_saved_document_is():
    text = "SUPREME COURT OF THE STATE OF NEW YORK COUNTY OF KINGS Index No. 512345/2026 " + "decision and order " * 20
    base = {"case": "Alpha Med v. Aetna", "court": "Supreme Court, Kings County", "level": "State Trial", "state": "ny",
            "docket": "512345/2026", "date": "2026-10-12", "kind": "ruling", "quote": "Index No. 512345/2026"}
    agreed = reader.identify(text, ask=lambda model, system, prompt: dict(base))
    assert agreed["agreed"] and agreed["state"] == "NY"
    calls = []
    other = lambda model, system, prompt: calls.append(model) or dict(base, date="2026-10-13" if model == reader.L2 else base["date"])  # noqa: E731
    assert not reader.identify(text, ask=other)["agreed"] and len(calls) == 2 * reader.ROUNDS
    assert not reader.identify(text, ask=lambda m, s, p: dict(base, quote="Index No. 999"))["agreed"]  # quote not in the text
    assert reader.identify("too short", ask=lambda *a: dict(base)) == {"agreed": False, "note": "The file has no readable text."}


def test_newly_added_cases_get_their_earlier_entries_once(monkeypatch):
    import daily_update as du
    hits = {"document_number": [{"docket_id": 5, "document_number": 1, "attachment_number": 1, "filepath_local": "recap/a.1.1.pdf"},
                                {"docket_id": 5, "document_number": 1, "filepath_local": "recap/a.1.0.pdf"}],
            "description": [{"docket_id": 5, "id": 77, "entry_date_filed": "2026-09-01", "filepath_local": "recap/a.9.0.pdf",
                             "description": "ORDER granting petition to confirm arbitration award"}]}
    monkeypatch.setattr(du, "search_all", lambda api, token, q: hits["document_number" if "document_number" in q else "description"])
    state = {"dockets": {"5": {**CASE, "docketId": "5", "doc": "", "history": False, "screened": False},
                         "6": {**CASE, "docketId": "6", "doc": "x", "screened": True}}}
    found = {"ended": []}
    got = du.catch_up(None, "", state, found)
    assert state["dockets"]["5"]["doc"].endswith("a.1.0.pdf") and state["dockets"]["5"]["history"] is True  # never an exhibit
    assert [c["key"] for c in got] == ["rd:77"] and got[0]["fetch"] == [("pdf", "https://storage.courtlistener.com/recap/a.9.0.pdf")]
    assert du.catch_up(None, "", state, found) == []  # only once


def test_the_whole_daily_run_hangs_together(tmp_path, monkeypatch):
    """main() end to end, with the searches and the readers faked: what is confirmed reaches found.json, the rest the
    day's To check file, and an answer from an earlier file is put on the dashboard."""
    import shutil
    import daily_update as du
    import find_updates as fu
    import newsletter
    import to_check as tc
    from conftest import REPORT
    if not REPORT.exists():
        pytest.skip("acceptance report not in reports\\")
    (tmp_path / "reports").mkdir()
    shutil.copy(REPORT, tmp_path / "reports" / REPORT.name)
    for name, value in (("ROOT", tmp_path), ("AUTO", tmp_path / "auto"), ("LOGS", tmp_path / "logs"),
                        ("FOUND", tmp_path / "auto" / data.AUTO_FILE), ("STATE", tmp_path / "auto" / "state.json"),
                        ("SUMMARY_SOURCES", tmp_path / "auto" / "summary_sources.json")):
        monkeypatch.setattr(du, name, value)
    monkeypatch.setattr(fu, "ROOT", tmp_path)
    monkeypatch.setattr(du, "check_setup", lambda: (None, "token", "", []))
    monkeypatch.setattr(du, "ruling_candidates", lambda *a: [
        {**CASE, "case": "Agreed v. Plan", "key": "o:1", "date": "2026-10-18", "fetch": [("pdf", "a")]},
        {**CASE, "case": "Disagreed v. Plan", "docket": "2:26-cv-00009", "key": "o:2", "date": "2026-10-18", "fetch": [("pdf", "b")]}])
    monkeypatch.setattr(du, "new_dockets", lambda *a: 0)
    monkeypatch.setattr(du, "georgetown_leads", lambda *a: 0)
    monkeypatch.setattr(du, "cms_guidance", lambda *a: 0)
    monkeypatch.setattr(du, "docket_updates", lambda *a: [])
    monkeypatch.setattr(du, "catch_up", lambda *a: [])
    monkeypatch.setattr(du, "new_dol", lambda *a: 0)
    monkeypatch.setattr(du, "state_search_step", lambda *a: None)  # its web searches are faked in test_state_search.py
    def summaries_fail(report):  # e.g. CourtListener's request limit: the run must still finish and keep its rows
        raise RuntimeError("HTTP 429 from www.courtlistener.com")
    monkeypatch.setattr(du, "write_summaries", summaries_fail)
    monkeypatch.setattr(du, "order_text", lambda api, token, c: "text")
    monkeypatch.setattr(newsletter, "run", lambda *a: "not due")
    monkeypatch.setattr(reader, "read", lambda text, case, *a: {"agreed": True, "relevant": True, "result": "Win", "final": True,
                                                                "holding": "Confirmed.", "rounds": 1} if case.startswith("Agreed")
                        else {"agreed": False, "rounds": 3, "note": "Could not agree after 3 rounds."})
    # An answer left on an earlier day's list.
    store = {}
    f, _ = tc.add(tmp_path, store, [tc.item("rd:9", "Ruling", "No free copy.", {**CASE, "case": "Answered v. Plan",
                                                                                 "docket": "2:26-cv-00008", "date": "2026-10-10"})],
                  dt.date(2026, 10, 18))
    (tmp_path / "auto").mkdir()
    tc.save(tmp_path, store)
    import openpyxl
    wb = openpyxl.load_workbook(f)
    wb["To check"]["J2"] = "Loss"
    wb.save(f)

    assert du.main([]) == 0
    found = json.loads((tmp_path / "auto" / data.AUTO_FILE).read_text(encoding="utf-8"))
    assert sorted((d["case"], d["result"], bool(d.get("byHand"))) for d in found["decisions"]) == [
        ("Agreed v. Plan", "Win", False), ("Answered v. Plan", "Loss", True)]
    wb = openpyxl.load_workbook(tmp_path / tc.FOLDER / tc.FILE)
    assert [r[0] for r in wb["To check"].iter_rows(min_row=2, values_only=True)] == ["Disagreed v. Plan"]  # today's row
    assert [r[0] for r in wb["Settled"].iter_rows(min_row=2, values_only=True)] == ["Answered v. Plan"]  # same file


def test_a_saved_file_the_readers_think_irrelevant_still_goes_to_a_person(monkeypatch):
    import daily_update  # noqa: F401  (puts app\tools on the path)
    import state_cases as sc
    monkeypatch.setattr(reader, "identify", lambda text: {"agreed": True, "case": "Clinic v. Plan", "court": "Court of Chancery",
                                                          "level": "State Trial", "state": "DE", "docket": "2025-0011",
                                                          "date": "2026-02-16", "kind": "ruling"})
    monkeypatch.setattr(reader, "read", lambda *a: {"agreed": True, "relevant": False, "result": "Pending", "final": False,
                                                    "holding": "", "reason": "It decides a payment claim.", "rounds": 1})
    what, adds, why, _ = sc.check_file("text", "rules", lambda c: "Plan")
    assert what == "to check" and adds == {} and "Leave out" in why
    monkeypatch.setattr(reader, "identify", lambda text: {"agreed": True, "case": "Clinic v. Plan", "court": "Court of Chancery",
                                                          "level": "State Trial", "state": "DE", "docket": "2025-0011",
                                                          "date": "2026-02-16", "kind": "complaint"})
    monkeypatch.setattr(reader, "screen", lambda text, case: False)
    assert sc.check_file("text", "rules", lambda c: "Plan")[0] == "to check"


def test_state_courts_get_a_wider_search_that_looks_back_once(monkeypatch):
    import daily_update as du
    import find_updates as fu
    calls = []

    def search(api, token, since, kind, pages, query=fu.CL_QUERY):
        calls.append((query == du.STATE_QUERY, since))
        if query != du.STATE_QUERY:
            return []
        return [{"court_id": "nysupctbroome", "court": "New York Supreme Court, Broome County", "court_jurisdiction": "ST",
                 "caseName": "Intraoperative Neurology Servs., PC v. Excellus BCBS", "docketNumber": "EFCA2025001234",
                 "dateFiled": "2026-08-24", "absolute_url": "/opinion/1/x/", "opinions": [{"id": 1}]},
                {"court_id": "nyed", "court": "District Court, E.D. New York", "caseName": "Clinic v. Aetna",
                 "docketNumber": "2:26-cv-00001", "dateFiled": "2026-08-20", "absolute_url": "/opinion/2/y/", "opinions": [{"id": 2}]}]
    monkeypatch.setattr(fu, "courtlistener", search)
    state, since = {}, dt.date(2026, 10, 16)
    got = du.ruling_candidates(None, "token", "", since, [], state)
    assert [c["case"] for c in got] == ["Intraoperative Neurology Servs., PC v. Excellus BCBS"]  # state courts only
    assert calls == [(False, since), (True, du.STATE_FROM)] and state["stateSearchDone"]  # the first run looks back to 2022
    calls.clear()
    du.ruling_candidates(None, "token", "", since, [], state)
    assert calls == [(False, since), (True, since)]


FEED = """<?xml version="1.0"?><rss xmlns:content="http://purl.org/rss/1.0/modules/content/"><channel>{}</channel></rss>"""
ITEM = """<item><title>{t}</title><link>https://example.org/{g}</link><guid>{g}</guid><pubDate>Mon, 21 Sep 2026 00:45:24 +0000</pubDate>
<content:encoded>&lt;p&gt;Docket {d}&lt;/p&gt;</content:encoded></item>"""


def test_new_cases_on_georgetowns_tracker_become_leads_for_a_person():
    import daily_update as du
    report = {"decisions": {"A": {"case": "Known Clinic v. Cigna", "docket": "2:25-cv-00100"}}, "pending": {}}
    old = ITEM.format(t="Old Clinic v. Aetna", g="1", d="2:24-cv-00001")
    state, asks = {}, []
    assert du.georgetown_leads(state, report, asks, fetch=lambda url: FEED.format(old)) == 0  # first run: only remembers
    assert asks == [] and state["georgetown"] == ["1"]
    items = old + "".join([ITEM.format(t="New Surgery PC v. Oxford Health Plans", g="2", d="1:26-cv-04321"),
                           ITEM.format(t="Known Clinic v. Cigna", g="3", d="2:25-cv-00100"),  # already on the dashboard
                           ITEM.format(t="Aetna Health Inc. v. Radiology Partners", g="4", d="3:26-cv-1")])  # insurer suing
    assert du.georgetown_leads(state, report, asks, fetch=lambda url: FEED.format(items)) == 1
    assert [(a["case"], a["docket"], a["level"], a["date"]) for a in asks] == [
        ("New Surgery PC v. Oxford Health Plans", "1:26-cv-04321", "Federal District", "2026-09-21")]
    asks.clear()
    assert du.georgetown_leads(state, report, asks, fetch=lambda url: FEED.format(items)) == 0  # never asked twice


def test_cases_that_are_clearly_not_enforcement_suits_are_settled_without_reading(monkeypatch):
    import daily_update as du
    import find_updates as fu
    assert fu.out_of_scope_nature("850 Securities/Commodities") and fu.out_of_scope_nature("410 Anti-Trust")
    assert fu.out_of_scope_nature("899 Other Statutes: Administrative Procedures Act/Review or Appeal of Agency Decision")
    assert not fu.out_of_scope_nature("896 Other Statutes: Arbitration") and not fu.out_of_scope_nature("Contract: Insurance")
    assert not fu.out_of_scope_nature("791 Labor: E.R.I.S.A.") and not fu.out_of_scope_nature("")
    assert "arbitration award" in fu.docket_facts("896 Other Statutes: Arbitration", "09:1 U.S. Arbitration Act")
    reads = []
    monkeypatch.setattr(reader, "screen", lambda text, case: reads.append(case) or None)
    monkeypatch.setattr(du, "pdf_text", lambda url: "text")
    base = {"court": "District Court, D. New Jersey", "level": "Federal District", "circuit": "3", "state": "NJ",
            "filed": "2026-06-23", "docket": "2:26-cv-1", "url": "", "payer": "", "doc": "x.pdf", "screened": False}
    state = {"dockets": {
        "1": {**base, "case": "Horizon Healthcare Services, Inc. v. M&D Capital Premier Billing, LLC", "docketId": "1"},
        "2": {**base, "case": "Bhagavan v. Nutex Health Inc.", "docketId": "2", "nature": "850 Securities/Commodities"},
        "3": {**base, "case": "CHS v. MultiPlan, Inc.", "docketId": "3", "nature": "410 Anti-Trust"},
        "4": {**base, "case": "Atlas Neurophysiological v. HealthComp, LLC", "docketId": "4",
              "nature": "896 Other Statutes: Arbitration", "cause": "09:1 U.S. Arbitration Act"}}}
    found, asks = {"pending": [{"docketId": "1", "case": "x"}]}, []
    du.screen_dockets(state, found, 10, asks)
    assert reads == ["Atlas Neurophysiological v. HealthComp, LLC"]  # only the possible enforcement suit is read
    assert [state["dockets"][i]["enforcement"] for i in "123"] == [False, False, False] and found["pending"] == []
    assert [a["case"] for a in asks] == ["Atlas Neurophysiological v. HealthComp, LLC"] and "arbitration award" in asks[0]["readers"]


def test_rows_the_daily_update_settles_itself_leave_the_list(tmp_path):
    import openpyxl
    import daily_update as du
    import to_check as tc
    store = {}
    row = lambda n, case, docket: tc.item(f"docket:{n}", "New case", "No free copy of the first filing.",  # noqa: E731
                                          {**CASE, "case": case, "docket": docket, "docketId": str(n), "date": "2026-09-01"})
    items = [row(1, "Ahuja Clinic v. Aetna", "2:25-cv-17642"), row(2, "Bhagavan v. Nutex Health Inc.", "4:25-cv-2"),
             row(3, "Rowe Plastic Surgery v. Horizon", "2:25-cv-3"), row(4, "Elite Neuro v. Cigna", "2:25-cv-4"),
             row(5, "Still Open v. Aetna", "2:25-cv-5")]
    f, _ = tc.add(tmp_path, store, items, dt.date(2026, 10, 19))
    state = {"dockets": {"1": {"enforcement": None}, "2": {"enforcement": False, "nature": "850 Securities/Commodities"},
                         "3": {"enforcement": None}, "4": {"enforcement": True}, "5": {"enforcement": None}}}
    found = {"decisions": [{**CASE, "docket": "2:25-cv-17642", "date": "2026-09-29", "result": "Loss", "holding": "Dismissed."}],
             "pending": [{"docketId": "4"}], "ended": [{"docket": "2:25-cv-3", "state": "NJ", "date": "2026-10-01",
                                                       "status": "Ended: voluntary dismissal or settlement"}]}
    assert du.auto_settle(store, state, found) == 4
    got = {k: (store[k]["answer"], store[k]["auto"]) for k in store if store[k].get("answer")}
    assert got == {"docket:1": ("Loss", True), "docket:2": ("Leave out", True), "docket:3": ("Leave out", True),
                   "docket:4": ("Pending case", True)}
    assert "850 Securities" in store["docket:2"]["note"] and "voluntary dismissal or settlement" in store["docket:3"]["note"]
    rows = tc.dashboard_rows(store)[0]  # the update's own Win/Loss stays the readers' row; its Pending case is rebuilt
    assert rows["decisions"] == [] and [(p["case"], p["byHand"], p["fromOrder"]) for p in rows["pending"]] == [
        ("Elite Neuro v. Cigna", False, False)]
    assert tc.tidy(tmp_path, store) == 4
    wb = openpyxl.load_workbook(f)
    assert [r[0].value for r in wb["To check"].iter_rows(min_row=2)] == ["Still Open v. Aetna"]
    assert {r[14].value for r in wb["Settled"].iter_rows(min_row=2)} == {"the daily update"}
    # A person disagrees on the Settled sheet: their answer wins from then on.
    ws = wb["Settled"]
    r = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "Bhagavan v. Nutex Health Inc.")
    ws.cell(r, 10).value = "Pending case"
    wb.save(f)
    assert tc.collect(tmp_path, store) == ["docket:2"] and "auto" not in store["docket:2"]
    assert [(p["case"], p["byHand"]) for p in tc.dashboard_rows(store)[0]["pending"]] == [
        ("Bhagavan v. Nutex Health Inc.", True), ("Elite Neuro v. Cigna", False)]


def test_cases_with_no_free_first_filing_are_looked_up_again_weekly():
    import daily_update as du
    store = {"docket:1": {"docketId": "1"}, "docket:2": {"docketId": "2"}, "docket:3": {"docketId": "3", "answer": "Win"}}
    state = {"dockets": {"1": {"doc": "", "enforcement": None, "screened": True, "history": True},
                         "2": {"doc": "x.pdf", "enforcement": None, "screened": True},
                         "3": {"doc": "", "enforcement": None, "screened": True}}}
    assert du.retry_missing_filings(store, state, dt.date(2026, 10, 20)) == 1
    assert state["dockets"]["1"] == {"doc": "", "enforcement": None, "screened": False, "history": False, "retried": "2026-10-20"}
    assert du.retry_missing_filings(store, state, dt.date(2026, 10, 25)) == 0  # not again within a week
    assert du.retry_missing_filings(store, state, dt.date(2026, 10, 27)) == 1


def test_bare_opinions_are_read_and_filings_to_read_go_first(monkeypatch):
    import check_dockets as cd
    import daily_update as du
    assert cd.entry_label("OPINION. Signed by Judge Susan D. Wigenton on 9/29/2026. (vm)") == "Ruling: opinion"
    assert "Ruling: opinion" in du.RULING_LABELS and "opinion" in cd.DECISIVE
    reads = []
    monkeypatch.setattr(reader, "screen", lambda text, case: reads.append(case) or None)
    monkeypatch.setattr(du, "pdf_text", lambda url: "text")
    base = {**CASE, "filed": "2026-10-01", "screened": False, "enforcement": None}
    state = {"dockets": {str(n): {**base, "case": f"Clinic {n} v. Aetna", "docketId": str(n), "doc": ""} for n in range(5)}}
    state["dockets"]["9"] = {**base, "case": "New Clinic v. Aetna", "docketId": "9", "doc": "x.pdf"}
    du.screen_dockets(state, {"pending": []}, 1, [])  # a limit of one: the case with a filing to read is not crowded out
    assert reads == ["New Clinic v. Aetna"] and all(d["screened"] for d in state["dockets"].values())


def test_an_open_to_check_case_gets_its_whole_docket_searched_once(monkeypatch):
    # Atlas v. HealthComp: on the To check list, but its stipulation of dismissal (Apr 2026) was before the look-back.
    import daily_update as du
    import to_check as tc
    queries = []

    class Api:
        def get(self, url, headers=None, body=None, cl=False):
            q = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)["q"][0]
            queries.append(q)
            hit = "70354791" in q and "entry_date_filed" not in q
            return {"results": [{"docket_id": 70354791, "description": "Stipulation of Dismissal",
                                 "entry_date_filed": "2026-04-01", "id": 5}] if hit else []}
    monkeypatch.setattr(du.cd, "resolve_ids", lambda *a: None)
    atlas = {**CASE, "case": "Atlas Neurophysiological Assessment, LLC v. HealthComp, LLC", "docket": "2:25-cv-02665",
             "state": "PA", "docketId": "70354791"}
    store = {"docket:70354791": tc.item("docket:70354791", "New case", "No free copy of the first filing.", atlas)}
    state = {"dockets": {"11": {**CASE, "case": "Clinic v. Aetna", "docketId": "11"}}}
    found = {"ended": [], "decisions": [], "pending": []}
    report = {"pending": {}, "decisions": {}}
    assert cd_entry_is_an_ending("Stipulation of Dismissal")
    du.docket_updates(Api(), "t", dt.date(2026, 8, 28), report, state, found, du.history_to_check(store, state))
    assert len(queries) == 2 and "docket_id:(11)" in queries[0] and "entry_date_filed:[2026-08-28" in queries[0]
    assert "docket_id:(70354791)" in queries[1] and "entry_date_filed" not in queries[1]  # the whole docket
    assert [(e["case"][:5], e["date"]) for e in found["ended"]] == [("Atlas", "2026-04-01")]
    assert du.history_to_check(store, state) == {}  # searched once: not again
    assert du.auto_settle(store, state, found) == 1 and store["docket:70354791"]["answer"] == tc.LEAVE


def test_new_jersey_cases_on_the_courts_stay_list_become_pending_cases(tmp_path):
    import daily_update as du
    import nj_stay
    import to_check as tc
    nj = lambda n, docket, **kw: tc.item(f"docket:{n}", "New case", "No free copy of the first filing.",  # noqa: E731
                                        {**CASE, "case": f"Clinic {n} v. Aetna", "docket": docket, "docketId": str(n),
                                         "filed": "2025-08-28", **kw})
    store = {i["key"]: i for i in (nj(1, "2:25-cv-14988"),                      # on Schedule A
                                   nj(2, "2:25-cv-17642"),                      # not on it (Ahuja's number)
                                   nj(3, "2:25-cv-15052"),                      # on it, but a ruling was read
                                   nj(4, "2:25-cv-15377"),                      # on it, but it ended
                                   nj(5, "2:25-cv-14988", court="District Court, S.D. New York", state="NY"))}  # same number, other court
    state = {"dockets": {str(n): {"case": f"Clinic {n} v. Aetna", "screened": True, "layers": 2, "enforcement": None}
                         for n in range(1, 6)}}
    found = {"decisions": [{**CASE, "docket": "2:25-cv-15052", "result": "Loss", "date": "2026-11-02", "holding": "Dismissed."}],
             "ended": [{**CASE, "docket": "2:25-cv-15377", "status": "Ended: voluntary dismissal or settlement", "date": "2026-02-01"}],
             "pending": []}
    assert du.auto_settle(store, state, found) == 3
    assert {k: i.get("answer") for k, i in store.items()} == {"docket:1": tc.PENDING, "docket:2": None, "docket:3": "Loss",
                                                               "docket:4": tc.LEAVE, "docket:5": None}
    assert "stay order of Jan 16, 2026" in store["docket:1"]["note"] and store["docket:1"]["auto"]
    assert [(p["docketId"], p["fromOrder"], p["status"]) for p in found["pending"]] == [("1", True, nj_stay.STATUS)]
    assert state["dockets"]["1"]["enforcement"] is True and du.waiting_dockets(state) == []  # never screened again
    assert du.retry_missing_filings(store, state) == 2 and "retried" not in state["dockets"]["1"]  # only the open rows 2 and 5
    # The next runs start by re-applying every answer: the stay-list case must stay on the dashboard, once, as it was
    # (2026-10-27: it was taken off on the second run, because only a person's answers were put back).
    for _ in range(2):
        du.apply_answers(store, state, found)
        assert [(p["docketId"], p["fromOrder"], p["byHand"], p["status"]) for p in found["pending"]] == [
            ("1", True, False, nj_stay.STATUS)]
    assert "byHand" not in state["dockets"]["1"]  # the update's own answer is not marked as a person's
    (tmp_path / data.AUTO_FILE).write_text(json.dumps(found), encoding="utf-8")
    decisions, pending = [], []
    data.merge_auto(tmp_path, decisions, pending, None, [])
    assert [(p["case"], p["fromOrder"], p["byHand"]) for p in pending] == [("Clinic 1 v. Aetna", True, False)]


def test_a_ruling_in_the_new_jersey_lead_case_is_called_out():
    import daily_update as du
    lead = {**CASE, "case": "East Coast Plastic Surgery PLLC PA v. Aetna", "docket": "2:25-cv-15052", "date": "2026-11-02"}
    other = {**CASE, "docket": "2:25-cv-15053", "date": "2026-11-02"}
    notes = du.lead_case_rulings([(lead, {"agreed": True, "result": "Loss"}, ""), (other, {"agreed": True, "result": "Loss"}, ""),
                                  (lead, {"agreed": False, "result": None}, "")])
    assert len(notes) == 1 and notes[0].startswith("LEAD CASE RULING: East Coast") and "329 New Jersey cases" in notes[0]


def cd_entry_is_an_ending(text):
    import check_dockets as cd
    return cd.entry_label(text) == "Ended: voluntary dismissal or settlement"


def test_new_guidance_on_cms_pages_goes_into_the_dol_actions():
    import daily_update as du
    sec = du.CMS_SECTION
    pages = {du.CMS_START[0]: '<a href="/initiatives/no-surprise-billing/idr">IDR</a> <a href="/nosurprises/x">consumers</a>'
                              '<a href="/initiatives/no-surprise-billing/x?y=1">query</a><a href="/files/document/old-faq.pdf">Old IDR FAQs</a>',
             sec + "idr": '<a href="/files/document/batching.pdf">Federal IDR Batching Notice</a>'
                                                    '<a href="/files/document/guide.pdf">IDR Gateway User Guide</a>'
                                                    '<a href="/files/document/fees.pdf">Administrative Fee FAQs</a>'
                                                    '<a href="/files/document/agents.pdf">Agent help</a>'}
    calls = []
    fetch = lambda url: calls.append(url) or pages.get(url, "")  # noqa: E731
    state, found = {}, {"dol": []}
    assert du.cms_guidance(state, [], found, fetch) == 0 and found["dol"] == []  # first run: only remembers
    assert sec + "x?y=1" not in calls and "https://www.cms.gov/nosurprises/x" not in calls  # robots.txt; consumer pages
    pages[sec + "idr"] += ('<a href="/files/document/ops.pdf">IDR Operations Update</a>'
                           '<a href="/files/document/manual.pdf">IDR Gateway User Guide v3</a>'
                           '<a href="/files/document/q3.xlsx">Federal IDR Supplemental Tables for 2026 Q3 (XLSX)</a>'
                           '<a href="/files/document/federal-idr-wildfire-attestation.pdf">published an attestation (PDF)</a>')
    pages[sec + "idr"] += '<a href="/files/document/set71.pdf">FAQs about CAA 2021 Implementation (Set 71)</a>'
    dates = lambda url: "2025-07-30" if url.endswith("ops.pdf") else ""  # noqa: E731
    assert du.cms_guidance(state, ["idr operations update 2", "faqs about caa 2021 and aca implementation part 71"],
                           found, fetch, dates) == 2  # Set 71 is the report's Part 71
    assert found["dol"][0]["issued"] == "2025-07-30" and found["dol"][1]["issued"] == dt.date.today().isoformat()
    assert [(a["title"], a["agency"], a["url"].split("/")[-1]) for a in found["dol"]] == [
        ("IDR Operations Update", "Centers for Medicare & Medicaid Services", "ops.pdf"),
        ("Federal IDR wildfire attestation", "Centers for Medicare & Medicaid Services", "federal-idr-wildfire-attestation.pdf")]
    assert du.cms_guidance(state, [], found, fetch, dates) == 0  # never added twice
    assert du.faq_number("FAQs About ACA Implementation Part 55 (PDF)") == "55" and du.faq_number("IDR Fee Notice") is None


def _daily_env(tmp_path, monkeypatch):
    """A daily update pointed at tmp_path, with every search and reader faked."""
    import shutil
    import daily_update as du
    import find_updates as fu
    import newsletter
    from conftest import REPORT
    if not REPORT.exists():
        pytest.skip("acceptance report not in reports\\")
    (tmp_path / "reports").mkdir()
    shutil.copy(REPORT, tmp_path / "reports" / REPORT.name)
    for name, value in (("ROOT", tmp_path), ("AUTO", tmp_path / "auto"), ("LOGS", tmp_path / "logs"),
                        ("FOUND", tmp_path / "auto" / data.AUTO_FILE), ("STATE", tmp_path / "auto" / "state.json"),
                        ("SUMMARY_SOURCES", tmp_path / "auto" / "summary_sources.json")):
        monkeypatch.setattr(du, name, value)
    monkeypatch.setattr(fu, "ROOT", tmp_path)
    monkeypatch.setattr(du, "check_setup", lambda: (None, "token", "", []))
    for step in ("new_dockets", "georgetown_leads", "cms_guidance", "new_dol"):
        monkeypatch.setattr(du, step, lambda *a: 0)
    for step in ("ruling_candidates", "docket_updates", "catch_up"):
        monkeypatch.setattr(du, step, lambda *a: [])
    monkeypatch.setattr(du, "write_summaries", lambda report: {})
    monkeypatch.setattr(newsletter, "run", lambda *a: "not due")
    return du


def test_a_crash_part_way_keeps_what_was_done(tmp_path, monkeypatch):
    import to_check as tc
    du = _daily_env(tmp_path, monkeypatch)

    def screen_then_crash(state, found, limit, asks, progress):
        found["pending"].append({**CASE, "docketId": "7"})
        asks.append(tc.item("docket:8", "New case", "The two readers could not tell.", {**CASE, "docketId": "8", "date": "2026-10-01"}))
        progress()
        raise RuntimeError("the network went away")
    monkeypatch.setattr(du, "waiting_dockets", lambda state: [1])
    monkeypatch.setattr(du, "screen_dockets", screen_then_crash)
    with pytest.raises(RuntimeError):
        du.main([])
    found = json.loads((tmp_path / "auto" / data.AUTO_FILE).read_text(encoding="utf-8"))
    assert [p["docketId"] for p in found["pending"]] == ["7"]  # the finds were saved
    assert "docket:8" in tc.load(tmp_path) and (tmp_path / tc.FOLDER / tc.FILE).exists()  # the row too
    assert not (tmp_path / "auto" / "daily_update.lock").exists()  # and the lock released


def test_a_to_check_file_open_in_excel_does_not_lose_rows(tmp_path, monkeypatch):
    import openpyxl
    import to_check as tc
    store, today = {}, dt.date(2026, 10, 20)
    tc.add(tmp_path, store, [tc.item("docket:1", "New case", "why", {**CASE, "docketId": "1"})], today)
    with monkeypatch.context() as m:  # open in Excel: nothing can be saved into it
        m.setattr(openpyxl.Workbook, "save", lambda self, filename: (_ for _ in ()).throw(PermissionError("open")))
        assert tc.add(tmp_path, store, [tc.item("docket:2", "New case", "why", {**CASE, "docketId": "2"})], today) == (None, 0)
        assert tc.unwritten(store) == 1 and tc.open_count(store) == 2  # kept, and counted as waiting
        assert tc.add(tmp_path, store, [], today) == (None, 0)  # still open on the next try
    f, n = tc.add(tmp_path, store, [], dt.date(2026, 10, 21))  # closed: the held-back row goes in at the bottom
    assert (f.name, n, tc.unwritten(store)) == (tc.FILE, 1, 0)
    assert [r[11] for r in openpyxl.load_workbook(f)["To check"].iter_rows(min_row=2, values_only=True)] == ["docket:1", "docket:2"]


def test_a_file_held_for_a_moment_is_still_saved(tmp_path, monkeypatch):
    real, held = os.replace, [2]

    def replace(src, dst):  # "Access is denied" twice, as when the file server is scanning the file
        if held[0]:
            held[0] -= 1
            raise PermissionError(5, "Access is denied")
        return real(src, dst)
    monkeypatch.setattr(data.os, "replace", replace)
    monkeypatch.setattr(data.time, "sleep", lambda s: None)
    import daily_update
    daily_update.save_json(tmp_path / "state.json", {"a": 1})
    assert json.loads((tmp_path / "state.json").read_text(encoding="utf-8")) == {"a": 1} and held == [0]
    held[0] = 99  # held for good: the error still stops the run (and its progress is saved by the caller)
    with pytest.raises(PermissionError):
        daily_update.save_json(tmp_path / "state.json", {"a": 2})


def test_a_data_file_that_cannot_be_read_is_never_saved_over(tmp_path, monkeypatch):
    p = tmp_path / "state.json"
    assert data.read_json(p, {"new": 1}) == {"new": 1}  # missing: a first run
    p.write_text('{"dockets": {"1": {}}}', encoding="utf-8")
    assert data.read_json(p, {}) == {"dockets": {"1": {}}} and (tmp_path / "state.json.bak").exists()  # good copy kept
    p.write_text('{"dockets": {"1": ', encoding="utf-8")  # half written
    assert data.read_json(p, {}) == {"dockets": {"1": {}}}  # the backup is used
    (tmp_path / "state.json.bak").write_text("garbage", encoding="utf-8")
    with pytest.raises(data.DataFileError):
        data.read_json(p, {})
    # The whole run stops without touching the file.
    du = _daily_env(tmp_path, monkeypatch)
    (tmp_path / "auto").mkdir()
    (tmp_path / "auto" / "state.json").write_text('{"half": ', encoding="utf-8")
    assert du.main([]) == 2
    assert (tmp_path / "auto" / "state.json").read_text(encoding="utf-8") == '{"half": '
    assert not (tmp_path / "auto" / data.AUTO_FILE).exists() and not (tmp_path / "auto" / "daily_update.lock").exists()


def test_only_one_daily_update_runs_at_a_time(tmp_path, monkeypatch):
    du = _daily_env(tmp_path, monkeypatch)
    (tmp_path / "auto").mkdir()
    assert du.acquire_lock() is None
    busy = du.acquire_lock()
    assert busy and "already running" in busy
    assert du.main([]) == 3  # a second run stops at once
    du.release_lock()
    old = {"user": "x", "computer": "y", "pid": 1, "started": (dt.datetime.now() - dt.timedelta(hours=9)).isoformat()}
    (tmp_path / "auto" / "daily_update.lock").write_text(json.dumps(old), encoding="utf-8")
    assert du.acquire_lock() is None  # a lock left by a run that died is taken over
    du.release_lock()
